from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException

from src.modules.commercial_core.schemas import (
    CatalogImportStatus,
    CatalogMatchStatus,
    CommercialCatalogImport,
    CommercialCatalogRow,
    CommercialCoreResponse,
    CommercialCoverage,
    CommercialEconomics,
    CommercialFeasibilityStatus,
    CommercialPositionMatch,
)
from src.modules.price_normalization.normalize import normalize_price, normalize_title
from src.modules.quote_comparison.position_matching import (
    ProcurementPosition,
    SupplierOfferCandidate,
    match_offer_to_position,
)
from src.shared.procurement_units import canonicalize_typed_position_unit

COMMERCIAL_CORE_CONTRACT_VERSION = "commercial-core-v1"
_MONEY = Decimal("0.01")

_HEADER_ALIASES: dict[str, set[str]] = {
    "sku": {"sku", "артикул", "article", "код", "код товара", "код номенклатуры", "номенклатурный номер"},
    "brand": {"бренд", "brand", "марка"},
    "manufacturer": {"производитель", "manufacturer", "изготовитель"},
    "title": {"наименование", "наименование товара", "товар", "название", "item", "item name", "product", "product name", "описание"},
    "unit": {"ед", "ед.", "единица", "единица измерения", "unit", "uom"},
    "price": {"цена", "цена за ед", "цена за единицу", "unit price", "price", "стоимость за ед"},
    "currency": {"валюта", "currency", "currency code"},
}


def _header_key(value: Any) -> str:
    text = str(value or "").casefold().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text, flags=re.IGNORECASE)
    return " ".join(text.split())


_ALIAS_LOOKUP = {
    _header_key(alias): field for field, aliases in _HEADER_ALIASES.items() for alias in aliases
}


@dataclass(frozen=True)
class _SheetRows:
    sheet_name: str
    rows: list[list[Any]]


@dataclass(frozen=True)
class _HeaderChoice:
    row_index: int
    mapping: dict[str, int]
    original_headers: dict[str, str]


class _LayoutNeedsReview(ValueError):
    pass


def _read_csv(content: bytes) -> list[_SheetRows]:
    last_error: Exception | None = None
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    if text is None:
        raise _LayoutNeedsReview(f"csv_decode_failed:{last_error}")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    return [_SheetRows(sheet_name="CSV", rows=rows)]


def _read_xlsx(content: bytes) -> list[_SheetRows]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result: list[_SheetRows] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        if any(any(cell not in (None, "") for cell in row) for row in rows):
            result.append(_SheetRows(sheet_name=sheet.title, rows=rows))
    return result


def _read_tabular(filename: str, content: bytes) -> list[_SheetRows]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _read_csv(content)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(content)
    raise _LayoutNeedsReview(f"unsupported_catalog_format:{suffix or 'none'}")


def _header_mapping(row: list[Any]) -> tuple[dict[str, int], dict[str, str], list[str]]:
    candidates: dict[str, list[int]] = {}
    originals: dict[str, str] = {}
    for index, value in enumerate(row):
        key = _header_key(value)
        field = _ALIAS_LOOKUP.get(key)
        if not field:
            continue
        candidates.setdefault(field, []).append(index)
        originals[field] = str(value or "").strip()
    ambiguous = [field for field, indexes in candidates.items() if len(indexes) > 1]
    mapping = {field: indexes[0] for field, indexes in candidates.items() if len(indexes) == 1}
    return mapping, originals, ambiguous


def _header_score(mapping: dict[str, int]) -> int:
    score = 0
    score += 5 if "title" in mapping else 0
    score += 4 if "price" in mapping else 0
    score += 3 if "sku" in mapping else 0
    score += 1 if "brand" in mapping else 0
    score += 1 if "manufacturer" in mapping else 0
    score += 1 if "unit" in mapping else 0
    score += 1 if "currency" in mapping else 0
    return score


def _choose_header(rows: list[list[Any]]) -> _HeaderChoice:
    ranked: list[tuple[int, int, dict[str, int], dict[str, str], list[str]]] = []
    for row_index, row in enumerate(rows[:20]):
        mapping, originals, ambiguous = _header_mapping(row)
        score = _header_score(mapping)
        if score:
            ranked.append((score, row_index, mapping, originals, ambiguous))
    if not ranked:
        raise _LayoutNeedsReview("header_not_detected")
    ranked.sort(key=lambda item: (-item[0], item[1]))
    best = ranked[0]
    if best[4]:
        raise _LayoutNeedsReview("ambiguous_columns:" + ",".join(sorted(best[4])))
    if "title" not in best[2] or ("price" not in best[2] and "sku" not in best[2]):
        raise _LayoutNeedsReview("required_columns_not_detected:title_and_price_or_sku")
    if len(ranked) > 1 and ranked[1][0] == best[0] and ranked[1][1] != best[1]:
        raise _LayoutNeedsReview("ambiguous_header_rows")
    return _HeaderChoice(best[1], best[2], best[3])


def _cell(row: list[Any], index: int | None) -> Any:
    return row[index] if index is not None and index < len(row) else None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split()).strip()
    return text or None


def _money(value: Decimal) -> float:
    return float(value.quantize(_MONEY, rounding=ROUND_HALF_UP))


def _catalog_row(
    *,
    filename: str,
    sheet_name: str,
    row_number: int,
    raw_row: list[Any],
    header_row: list[Any],
    choice: _HeaderChoice,
    default_currency: str | None,
) -> CommercialCatalogRow | None:
    title = _text(_cell(raw_row, choice.mapping.get("title")))
    if not title:
        return None
    warnings: list[str] = []
    parsed_price = normalize_price(_cell(raw_row, choice.mapping.get("price")))
    if parsed_price is not None and parsed_price < 0:
        parsed_price = None
        warnings.append("negative_price_rejected")
    if parsed_price is None:
        warnings.append("price_unknown")
    currency = _text(_cell(raw_row, choice.mapping.get("currency")))
    if currency:
        currency = currency.upper()
    elif default_currency:
        currency = default_currency.strip().upper() or None
        if currency:
            warnings.append("currency_from_operator_default")
    else:
        warnings.append("currency_unknown")
    unit = canonicalize_typed_position_unit(_text(_cell(raw_row, choice.mapping.get("unit"))))
    mapped_indexes = set(choice.mapping.values())
    characteristics: dict[str, Any] = {}
    for index, header in enumerate(header_row):
        if index in mapped_indexes:
            continue
        header_text = _text(header)
        value = _cell(raw_row, index)
        if header_text and value not in (None, ""):
            characteristics[header_text] = value
    safe_sheet = re.sub(r"[^0-9A-Za-zА-Яа-я_-]+", "-", sheet_name).strip("-") or "sheet"
    return CommercialCatalogRow(
        row_id=f"catalog:{safe_sheet}:{row_number}",
        source_file=Path(filename).name,
        sheet_name=sheet_name,
        row_number=row_number,
        sku=_text(_cell(raw_row, choice.mapping.get("sku"))),
        brand=_text(_cell(raw_row, choice.mapping.get("brand"))),
        manufacturer=_text(_cell(raw_row, choice.mapping.get("manufacturer"))),
        title=title,
        unit=unit,
        price=_money(parsed_price) if parsed_price is not None else None,
        currency=currency,
        characteristics=characteristics,
        warnings=warnings,
    )


def import_catalog(
    filename: str,
    content: bytes,
    *,
    default_currency: str | None = None,
) -> CommercialCatalogImport:
    source_sha256 = hashlib.sha256(content).hexdigest()
    try:
        sheets = _read_tabular(filename, content)
    except (_LayoutNeedsReview, ValueError, OSError, BadZipFile, InvalidFileException) as exc:
        return CommercialCatalogImport(
            status=CatalogImportStatus.NEEDS_REVIEW,
            source_file=Path(filename).name,
            source_sha256=source_sha256,
            unknowns=[str(exc)],
        )
    rows: list[CommercialCatalogRow] = []
    detected_columns: dict[str, str] = {}
    warnings: list[str] = []
    unknowns: list[str] = []
    for sheet in sheets:
        try:
            choice = _choose_header(sheet.rows)
        except _LayoutNeedsReview as exc:
            unknowns.append(f"{sheet.sheet_name}:{exc}")
            continue
        header_row = sheet.rows[choice.row_index]
        for field, header in choice.original_headers.items():
            detected_columns.setdefault(field, header)
        for zero_based_index, raw_row in enumerate(sheet.rows[choice.row_index + 1 :], start=choice.row_index + 1):
            item = _catalog_row(
                filename=filename,
                sheet_name=sheet.sheet_name,
                row_number=zero_based_index + 1,
                raw_row=raw_row,
                header_row=header_row,
                choice=choice,
                default_currency=default_currency,
            )
            if item is not None:
                rows.append(item)
    if not rows:
        return CommercialCatalogImport(
            status=CatalogImportStatus.NEEDS_REVIEW,
            source_file=Path(filename).name,
            source_sha256=source_sha256,
            detected_columns=detected_columns,
            warnings=warnings,
            unknowns=unknowns or ["no_catalog_rows"],
        )
    normalized_skus = [re.sub(r"[^0-9a-zа-я]+", "", (row.sku or "").casefold()) for row in rows if row.sku]
    duplicate_skus = sorted(key for key, count in Counter(normalized_skus).items() if key and count > 1)
    if duplicate_skus:
        warnings.append("duplicate_sku:" + ",".join(duplicate_skus))
    status = CatalogImportStatus.NEEDS_REVIEW if unknowns else CatalogImportStatus.READY
    return CommercialCatalogImport(
        status=status,
        source_file=Path(filename).name,
        source_sha256=source_sha256,
        rows=rows,
        detected_columns=detected_columns,
        warnings=warnings,
        unknowns=unknowns,
    )


def _unit_key(value: str | None) -> str | None:
    canonical = canonicalize_typed_position_unit(value)
    if not canonical:
        return None
    aliases = {
        "штука": "шт", "штук": "шт", "шт.": "шт",
        "метр": "м", "метров": "м", "м.": "м",
        "килограмм": "кг", "килограммов": "кг", "кг.": "кг",
        "комплект": "компл", "компл.": "компл",
    }
    return aliases.get(canonical.casefold(), canonical.casefold())


def _position_payload(item: dict[str, Any], index: int) -> ProcurementPosition | None:
    name = _text(item.get("official_name") or item.get("display_name") or item.get("normalized_name") or item.get("name"))
    if not name:
        return None
    return ProcurementPosition(
        position_id=str(item.get("stable_item_id") or item.get("position_id") or f"line-{index}"),
        item_name=name,
        quantity=item.get("quantity"),
        unit=_text(item.get("unit_normalized") or item.get("unit_original") or item.get("unit")),
        manufacturer=_text(item.get("manufacturer")),
        brand=_text(item.get("brand")),
        model=_text(item.get("model")),
        article=_text(item.get("article") or item.get("sku")),
    )


def _catalog_offer(row: CommercialCatalogRow) -> SupplierOfferCandidate:
    return SupplierOfferCandidate(
        offer_id=row.row_id,
        supplier_label="customer_catalog",
        item_name=row.title,
        source_type="commercial_quote",
        source_ref=f"{row.source_file}:{row.sheet_name}:row:{row.row_number}",
        currency_code=row.currency or "UNKNOWN",
        unit_price=row.price,
        manufacturer=row.manufacturer,
        brand=row.brand,
        article=row.sku,
    )


def _normalized_identifier(value: str | None) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", (value or "").casefold())


def _match_status(
    position: ProcurementPosition,
    row: CommercialCatalogRow,
    score: float,
    *,
    ambiguous: bool,
) -> tuple[CatalogMatchStatus, list[str]]:
    rationale: list[str] = []
    position_unit = _unit_key(position.unit)
    row_unit = _unit_key(row.unit)
    if position_unit and row_unit and position_unit != row_unit:
        return CatalogMatchStatus.UNCERTAIN, [f"unit_conflict:{position_unit}!={row_unit}"]
    if position_unit and not row_unit:
        return CatalogMatchStatus.UNCERTAIN, ["catalog_unit_unknown"]
    if row_unit and not position_unit:
        return CatalogMatchStatus.UNCERTAIN, ["tender_unit_unknown"]
    if ambiguous:
        return CatalogMatchStatus.UNCERTAIN, ["multiple_close_catalog_candidates"]
    article_match = bool(
        _normalized_identifier(position.article)
        and _normalized_identifier(position.article) == _normalized_identifier(row.sku)
    )
    same_title = normalize_title(position.item_name) == normalize_title(row.title)
    brand_or_mfr_match = bool(
        (_normalized_identifier(position.brand) and _normalized_identifier(position.brand) == _normalized_identifier(row.brand))
        or (_normalized_identifier(position.manufacturer) and _normalized_identifier(position.manufacturer) == _normalized_identifier(row.manufacturer))
    )
    if article_match:
        rationale.append("article_exact")
        return CatalogMatchStatus.EXACT, rationale
    if same_title and (brand_or_mfr_match or not any([position.brand, position.manufacturer, row.brand, row.manufacturer])):
        rationale.append("normalized_title_exact")
        if brand_or_mfr_match:
            rationale.append("brand_or_manufacturer_exact")
        return CatalogMatchStatus.EXACT, rationale
    if score >= 0.42:
        return CatalogMatchStatus.LIKELY_ANALOG, [f"generic_match_score:{score:.4f}"]
    if score >= 0.30:
        return CatalogMatchStatus.PARTIAL, [f"generic_match_score:{score:.4f}"]
    if score >= 0.12:
        return CatalogMatchStatus.UNCERTAIN, [f"weak_match_score:{score:.4f}"]
    return CatalogMatchStatus.NO_MATCH, [f"match_score_below_floor:{score:.4f}"]


def _match_position(position: ProcurementPosition, rows: list[CommercialCatalogRow]) -> CommercialPositionMatch:
    ranked: list[tuple[float, CommercialCatalogRow, list[str]]] = []
    for row in rows:
        match = match_offer_to_position(position, _catalog_offer(row), match_threshold=0.0)
        ranked.append((match.match_score, row, match.match_reasons))
    ranked.sort(key=lambda item: (-item[0], item[1].row_id))
    if not ranked:
        return CommercialPositionMatch(
            position_id=position.position_id,
            tender_name=position.item_name,
            tender_quantity=float(normalize_price(position.quantity)) if normalize_price(position.quantity) is not None else None,
            tender_unit=position.unit,
            status=CatalogMatchStatus.NO_MATCH,
            rationale=["catalog_empty"],
        )
    top_score, top_row, generic_reasons = ranked[0]
    close = [item for item in ranked if item[0] >= 0.30 and abs(item[0] - top_score) <= 0.02]
    ambiguous = len(close) > 1
    if "article_conflict" in generic_reasons:
        status, rationale = CatalogMatchStatus.NO_MATCH, ["article_conflict"]
    else:
        status, rationale = _match_status(position, top_row, top_score, ambiguous=ambiguous)
    selected = status in {CatalogMatchStatus.EXACT, CatalogMatchStatus.LIKELY_ANALOG, CatalogMatchStatus.PARTIAL}
    evidence = []
    if selected or status == CatalogMatchStatus.UNCERTAIN:
        evidence = [{
            "source_file": top_row.source_file,
            "sheet": top_row.sheet_name,
            "row": top_row.row_number,
            "catalog_row_id": top_row.row_id,
        }]
    quantity = normalize_price(position.quantity)
    return CommercialPositionMatch(
        position_id=position.position_id,
        tender_name=position.item_name,
        tender_quantity=float(quantity) if quantity is not None else None,
        tender_unit=position.unit,
        status=status,
        catalog_row_id=top_row.row_id if selected else None,
        catalog_title=top_row.title if selected else None,
        catalog_unit=top_row.unit if selected else None,
        catalog_unit_price=top_row.price if selected else None,
        currency=top_row.currency if selected else None,
        match_score=top_score,
        rationale=rationale + generic_reasons,
        evidence=evidence,
    )


def _float_number(value: Any) -> float | None:
    parsed = normalize_price(value)
    return float(parsed) if parsed is not None else None


def _coverage(matches: list[CommercialPositionMatch]) -> CommercialCoverage:
    counts = Counter(match.status for match in matches)
    total = len(matches)
    matched = sum(counts[status] for status in (CatalogMatchStatus.EXACT, CatalogMatchStatus.LIKELY_ANALOG, CatalogMatchStatus.PARTIAL))
    costed = sum(
        1 for match in matches
        if match.status in {CatalogMatchStatus.EXACT, CatalogMatchStatus.LIKELY_ANALOG}
        and match.catalog_unit_price is not None
        and match.tender_quantity is not None
        and not any(reason.startswith("unit_conflict") for reason in match.rationale)
    )
    return CommercialCoverage(
        total_positions=total,
        exact=counts[CatalogMatchStatus.EXACT],
        likely_analog=counts[CatalogMatchStatus.LIKELY_ANALOG],
        partial=counts[CatalogMatchStatus.PARTIAL],
        uncertain=counts[CatalogMatchStatus.UNCERTAIN],
        no_match=counts[CatalogMatchStatus.NO_MATCH],
        matched_coverage_ratio=round(matched / total, 4) if total else None,
        costed_coverage_ratio=round(costed / total, 4) if total else None,
    )


def _economics(
    model: dict[str, Any],
    matches: list[CommercialPositionMatch],
    *,
    target_bid_amount: float | None,
) -> CommercialEconomics:
    costed: list[CommercialPositionMatch] = []
    unknown: list[str] = []
    currencies: set[str] = set()
    for match in matches:
        if match.status not in {CatalogMatchStatus.EXACT, CatalogMatchStatus.LIKELY_ANALOG}:
            unknown.append(match.position_id)
            continue
        if match.catalog_unit_price is None or match.tender_quantity is None or not match.currency:
            unknown.append(match.position_id)
            continue
        tender_unit = _unit_key(match.tender_unit)
        catalog_unit = _unit_key(match.catalog_unit)
        if not tender_unit or not catalog_unit or tender_unit != catalog_unit:
            unknown.append(match.position_id)
            continue
        costed.append(match)
        currencies.add(match.currency.upper())
    if len(currencies) != 1:
        return CommercialEconomics(
            status="UNKNOWN",
            costed_position_ids=[item.position_id for item in costed],
            unknown_cost_position_ids=sorted(set(unknown + [item.position_id for item in costed])),
            target_bid_amount=target_bid_amount,
            nmck_ceiling=_float_number(model.get("nmck")),
            warnings=["catalog_currency_not_uniform_or_unknown"],
        )
    currency = next(iter(currencies))
    if unknown or len(costed) != len(matches):
        return CommercialEconomics(
            status="PARTIAL",
            currency=currency,
            costed_position_ids=[item.position_id for item in costed],
            unknown_cost_position_ids=sorted(set(unknown)),
            target_bid_amount=target_bid_amount,
            nmck_ceiling=_float_number(model.get("nmck")),
            warnings=["full_catalog_cost_not_supported"],
        )
    total = Decimal(0)
    for match in costed:
        total += Decimal(str(match.catalog_unit_price)) * Decimal(str(match.tender_quantity))
    known_catalog_cost = _money(total)
    economics = CommercialEconomics(
        status="KNOWN",
        currency=currency,
        known_catalog_cost=known_catalog_cost,
        costed_position_ids=[item.position_id for item in costed],
        unknown_cost_position_ids=[],
        target_bid_amount=target_bid_amount,
        nmck_ceiling=_float_number(model.get("nmck")),
    )
    if target_bid_amount is not None:
        margin = Decimal(str(target_bid_amount)) - total
        economics.gross_margin_amount = _money(margin)
        if Decimal(str(target_bid_amount)) != 0:
            economics.gross_margin_percent = round(float(margin / Decimal(str(target_bid_amount)) * Decimal(100)), 2)
    nmck = normalize_price(model.get("nmck"))
    report_currency = str(model.get("currency") or "").upper()
    if nmck is not None and report_currency == currency:
        headroom = nmck - total
        economics.nmck_headroom_amount = _money(headroom)
        if nmck != 0:
            economics.nmck_headroom_percent = round(float(headroom / nmck * Decimal(100)), 2)
    elif nmck is not None:
        economics.warnings.append("nmck_currency_mismatch_or_unknown")
    return economics


def build_commercial_core(
    model: dict[str, Any],
    *,
    catalog_filename: str,
    catalog_content: bytes,
    default_currency: str | None = None,
    target_bid_amount: float | None = None,
) -> CommercialCoreResponse:
    catalog = import_catalog(catalog_filename, catalog_content, default_currency=default_currency)
    positions = [
        position for index, item in enumerate(model.get("line_items", []) or [], start=1)
        if isinstance(item, dict) and (position := _position_payload(item, index)) is not None
    ]
    matches = [_match_position(position, catalog.rows) for position in positions] if catalog.rows else []
    coverage = _coverage(matches)
    economics = _economics(model, matches, target_bid_amount=target_bid_amount) if matches else CommercialEconomics(
        status="UNKNOWN",
        target_bid_amount=target_bid_amount,
        nmck_ceiling=_float_number(model.get("nmck")),
        warnings=["no_tender_positions_or_catalog_matches"],
    )
    rationale: list[str] = []
    if catalog.status != CatalogImportStatus.READY:
        feasibility = CommercialFeasibilityStatus.NEEDS_REVIEW
        rationale.append("Каталог импортирован неоднозначно; требуется проверка структуры/колонок.")
    elif not positions:
        feasibility = CommercialFeasibilityStatus.NEEDS_REVIEW
        rationale.append("В каноническом отчёте нет подтверждённых товарных позиций для коммерческого матчинга.")
    elif any(match.status != CatalogMatchStatus.EXACT for match in matches):
        feasibility = CommercialFeasibilityStatus.NEEDS_REVIEW
        rationale.append("Не все позиции имеют однозначный EXACT-матч; аналоги/частичные/неопределённые совпадения требуют решения человека.")
    elif economics.status != "KNOWN" or economics.known_catalog_cost is None:
        feasibility = CommercialFeasibilityStatus.NEEDS_REVIEW
        rationale.append("Полная себестоимость по каталогу не подтверждена из-за отсутствующих цены/количества/валюты.")
    else:
        limit = target_bid_amount if target_bid_amount is not None else economics.nmck_ceiling
        if limit is not None and economics.known_catalog_cost > limit:
            feasibility = CommercialFeasibilityStatus.NOT_FEASIBLE_AT_CURRENT_CATALOG_PRICE
            rationale.append("Подтверждённая стоимость по текущему каталогу превышает заданную цену предложения/НМЦК.")
        else:
            feasibility = CommercialFeasibilityStatus.FEASIBLE
            rationale.append("Все позиции имеют однозначный EXACT-матч и подтверждённую стоимость по каталогу.")
    next_action = {
        CommercialFeasibilityStatus.FEASIBLE: "Передать расчёт на человеческую коммерческую проверку и определить фактическую цену предложения.",
        CommercialFeasibilityStatus.NOT_FEASIBLE_AT_CURRENT_CATALOG_PRICE: "Проверить скидки, допустимые замены и альтернативных поставщиков; решение об участии принимает человек.",
        CommercialFeasibilityStatus.NEEDS_REVIEW: "Проверить отмеченные позиции/данные каталога и повторить коммерческий расчёт до решения об участии.",
    }[feasibility]
    return CommercialCoreResponse(
        contract_version=COMMERCIAL_CORE_CONTRACT_VERSION,
        catalog=catalog,
        matches=matches,
        coverage=coverage,
        economics=economics,
        feasibility_status=feasibility,
        rationale=rationale,
        next_action=next_action,
        human_control_required=True,
        external_action_allowed=False,
        safety={
            "bid_submission_allowed": False,
            "rfq_or_invitation_allowed": False,
            "supplier_contact_allowed": False,
            "purchase_allowed": False,
            "signing_allowed": False,
            "external_commercial_effect_allowed": False,
        },
    )
