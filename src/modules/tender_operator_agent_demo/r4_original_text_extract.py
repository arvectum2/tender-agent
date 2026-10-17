"""Read-only PDF/DOCX/ZIP source extraction, no legal text interpretation.

For scanned PDFs or unsupported spreadsheets, fail closed: the model must
never fabricate text or silently skip relevant source material.
"""
from __future__ import annotations

import io
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def _pdf_text(data: bytes) -> str:
    with tempfile.TemporaryDirectory() as temporary:
        p = Path(temporary) / "source.pdf"
        p.write_bytes(data)
        run = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", str(p), "-"],
            capture_output=True, check=True, timeout=40,
        )
    text = run.stdout.decode("utf8", errors="strict")
    if len(text.strip()) < 80:
        raise ValueError("PDF source requires separate visual/OCR review")
    return text


def _xlsx_text(data: bytes) -> str:
    """Read static XLSX cells for citation checks; reject formulas and large files."""
    if len(data) > 10_000_000:
        raise ValueError("XLSX file exceeds safe size")
    with zipfile.ZipFile(io.BytesIO(data)) as container:
        entries = container.infolist()
        if (
            len(entries) > 500
            or sum(item.file_size for item in entries) > 40_000_000
            or any(item.file_size > 20_000_000 or item.flag_bits & 1 for item in entries)
        ):
            raise ValueError("Unbounded or encrypted XLSX")
        if any(
            "externalLinks/" in item.filename or item.filename.startswith("xl/embeddings/")
            for item in entries
        ):
            raise ValueError("External-link or embedded XLSX source requires human review")
    workbook = load_workbook(
        io.BytesIO(data), read_only=True, data_only=False, keep_links=False,
    )
    try:
        lines = []
        for sheet in workbook.worksheets:
            rows, columns = sheet.max_row or 0, sheet.max_column or 0
            if rows * columns > 250_000 or rows > 100_000 or columns > 200:
                raise ValueError("Oversized XLSX sheet requires independent review")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        raise ValueError("Formula value is not a literal source fact")
                    if cell.data_type == "e":
                        raise ValueError("XLSX error cell requires independent review")
                    text = " ".join(str(cell.value).split())
                    if text:
                        lines.append(f"SHEET-{sheet.title} CELL-{cell.coordinate}: {text}")
        if not lines:
            raise ValueError("No literal XLSX source cells")
        return chr(10).join(lines)
    finally:
        workbook.close()


def verified_pdf_citation_pages(raw: bytes, filename: str) -> list[str] | None:
    """Extract PDF page texts for strictly localized citation verification.

    Preserve empty pages: removing one would shift all following citations.
    Layout extraction remains incomplete for scans/tables, so human legal
    review is still necessary even when a page quote is textually verified.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".zip":
        if len(raw) > 30_000_000:
            raise ValueError("Original ZIP too large for PDF page evidence")
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = [info for info in archive.infolist() if not info.is_dir()]
            if len(entries) != 1 or entries[0].file_size > 20_000_000:
                raise ValueError("Ambiguous/oversize nested original for PDF page proof")
            info = entries[0]
            if info.flag_bits & 1 or Path(info.filename).is_absolute() or ".." in Path(info.filename).parts:
                raise ValueError("Unsafe nested PDF original")
            if Path(info.filename).suffix.lower() != ".pdf":
                return None
            return verified_pdf_citation_pages(archive.read(info), info.filename)
    if suffix != ".pdf":
        return None
    pages = _pdf_text(raw).split("\f")
    if pages and not pages[-1].strip():
        pages.pop()
    if not pages or len(pages) > 1_500:
        raise ValueError("Unverifiable PDF page count")
    return pages


def verified_document_text(raw: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _pdf_text(raw)
    if suffix == ".xlsx":
        return _xlsx_text(raw)
    if suffix == ".docx":
        doc = Document(io.BytesIO(raw))
        lines = [
            f"PARAGRAPH-{index:04d}: {part.text}"
            for index, part in enumerate(doc.paragraphs, 1)
            if part.text.strip()
        ]
        for table_id, table in enumerate(doc.tables, 1):
            for row_id, row in enumerate(table.rows, 1):
                for cell_id, cell in enumerate(row.cells, 1):
                    if cell.text.strip():
                        lines.append(
                            f"TABLE-{table_id:04d} ROW-{row_id:04d} "
                            f"CELL-{cell_id:04d}: {cell.text}"
                        )
        if sum(len(line) for line in lines) < 80:
            raise ValueError("DOCX source insufficient for text comparison")
        return "\n".join(lines)
    if suffix != ".zip":
        raise ValueError("Unsupported official original document format")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = [info for info in archive.infolist() if not info.is_dir()]
        if len(entries) != 1 or entries[0].file_size > 20_000_000:
            raise ValueError("Ambiguous/oversize official nested document ZIP")
        info = entries[0]
        inner_path = Path(info.filename)
        if info.flag_bits & 1 or inner_path.is_absolute() or ".." in inner_path.parts:
            raise ValueError("Unsafe ZIP inner original document")
        if inner_path.suffix.lower() not in {".docx", ".pdf", ".xlsx"}:
            raise ValueError("Unsupported inner original source document")
        return verified_document_text(archive.read(info), info.filename)
