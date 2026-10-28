"""Prepare source-bound technical extracts for ChatGPT review, without a local LLM.

ChatGPT in this conversation is the reasoning engine. This module emits a
bounded review packet for an authorized caller to submit to ChatGPT, and
validates a structured human/model review before promotion.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml

ALLOWED = {".pdf", ".txt", ".doc", ".docx", ".xlsx"}
MAX_FILE = 20 * 1024 * 1024
MAX_TEXT = 30000
MAX_FILES = 24


def _extract(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix not in ALLOWED or path.is_symlink() or not path.is_file():
        raise ValueError("Unsupported or unsafe technical document")
    if path.stat().st_size > MAX_FILE:
        raise ValueError("Technical document exceeds budget")
    if suffix == ".pdf":
        proc = subprocess.run(["/opt/homebrew/bin/pdftotext", "-layout", str(path), "-"],
                              capture_output=True, timeout=25, check=False)
        if proc.returncode:
            return "", "pdf_text_extraction_failed"
        return proc.stdout.decode("utf-8", errors="replace")[:MAX_TEXT], "native_pdf_text"
    if suffix == ".doc":
        # Legacy binary Word originals are common in EIS 44/223. Preserve the
        # original file SHA; extraction is bounded, read-only, and macro-free.
        if path.open("rb").read(8) != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            return "", "invalid_legacy_doc_container"
        if sys.platform == "darwin":
            command = ["/usr/bin/textutil", "-convert", "txt", "-stdout", str(path)]
        else:
            binary = shutil.which("antiword")
            if binary is None:
                return "", "legacy_doc_extractor_unavailable"
            command = [binary, str(path)]
        try:
            proc = subprocess.run(command, capture_output=True, timeout=25, check=False)
        except (OSError, subprocess.TimeoutExpired):
            return "", "legacy_doc_extraction_failed"
        if proc.returncode:
            return "", "legacy_doc_extraction_failed"
        return proc.stdout.decode("utf-8", errors="replace")[:MAX_TEXT], "legacy_doc_text"
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT], "text"
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            chunks.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        return "\n".join(chunks)[:MAX_TEXT], "docx"
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for sheet in book.worksheets[:12]:
        chunks.append(f"[SHEET {sheet.title}]")
        for row in sheet.iter_rows(max_row=2500, max_col=50, values_only=True):
            chunks.append(" | ".join(str(c)[:250] if c is not None else "" for c in row))
            if sum(map(len, chunks)) >= MAX_TEXT:
                break
    book.close()
    return "\n".join(chunks)[:MAX_TEXT], "xlsx"


def prepare_packet(*, document_dir: Path, ontology_path: Path,
                   registry_number: str, max_files: int = 8) -> dict:
    if not re.fullmatch(r"[0-9]{11}|[0-9]{19,20}", registry_number):
        raise ValueError("Invalid registry number")
    if not 1 <= max_files <= MAX_FILES or document_dir.is_symlink():
        raise ValueError("Invalid document scope")
    ontology = yaml.safe_load(ontology_path.read_text(encoding="utf-8"))
    terms = [(family["id"], word) for family in ontology["families"]
             for word in family.get("positive", []) if len(word) >= 3]
    docs = []
    extraction_errors = []
    for path in sorted(document_dir.iterdir()):
        if len(docs) >= max_files:
            break
        if path.suffix.lower() not in ALLOWED:
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        try:
            content, method = _extract(path)
        except zipfile.BadZipFile:
            # Some publicly uploaded .docx/.xlsx files are mislabeled RTF
            # or corrupt ZIP bytes. Keep their original SHA, do not fabricate
            # extracted text, and continue the other source documents.
            content, method = "", "invalid_office_archive_container"
            extraction_errors.append({
                "sha256": digest, "filename": path.name,
                "error_type": "BadZipFile",
            })
        hits = [{"family": family, "term": word} for family, word in terms
                if word.casefold() in content.casefold()]
        docs.append({"filename": path.name, "sha256": digest, "method": method,
                     "text_truncated": len(content) >= MAX_TEXT,
                     "text": content, "lexical_hints_only": hits[:100]})
    return {
        "schema_version": "pq-chatgpt-review-packet/1",
        "registry_number": registry_number,
        "documents": docs,
        "extraction_errors": extraction_errors,
        "review_instructions": (
            "Act as a Russian 6-35 kV Power Quality procurement analyst. "
            "Check original document quotations and precise equipment scope. "
            "Classify UKRM, SVC/TSC, SVG/STATCOM, harmonic filters and substation EPC "
            "only when text supports it; extract kV, kvar, quantities, prices and "
            "buyer/supplier roles with SHA and quotations. "
            "Do not invent winners, contract awards, manufacturer, NMC or totals. "
            "Distinguish design, supply, installation and signed contract. "
            "Return review with evidence and uncertainties."
        ),
        "llm_executed": False,
        "market_total_rub": None,
    }


def save_packet(packet: dict, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(packet, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")


def validate_review(*, packet: dict, review: dict) -> dict:
    """Accept ChatGPT's classification only with exact source-backed quotations."""
    if review.get("registry_number") != packet.get("registry_number"):
        raise ValueError("Review registry mismatch")
    classification = review.get("classification")
    if classification not in ("relevant", "irrelevant", "uncertain"):
        raise ValueError("Invalid review classification")
    evidence = review.get("evidence")
    if not isinstance(evidence, list) or not evidence:
        raise ValueError("Review requires evidence")
    sources = {d["sha256"]: d["text"] for d in packet["documents"]}
    for item in evidence:
        sha = item.get("sha256")
        quote = item.get("quote")
        if sha not in sources or not isinstance(quote, str) or len(quote) < 12:
            raise ValueError("Unbound review evidence")
        if " ".join(quote.split()) not in " ".join(sources[sha].split()):
            raise ValueError("Review quote absent from original text")
    return {
        "schema_version": "pq-chatgpt-source-review/1",
        "registry_number": packet["registry_number"],
        "classification": classification,
        "evidence": evidence,
        "reasoning": str(review.get("reasoning", ""))[:1500],
        "signed_contract_verified": False,
        "winner_verified": False,
        "market_total_rub": None,
        "ocr_requires_visual_verification": any(d.get("method") == "ocr_unverified" for d in packet["documents"]),
    }
