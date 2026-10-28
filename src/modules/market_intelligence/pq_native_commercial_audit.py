"""Cross-check Watchdog's separately stored first-party commercial ETP originals.

The Watchdog report describes human-visual PDF readings. This auditor only
promotes SHA-verified first-party cards/files, and rechecks native DOCX quotes
for 35 kV owner-furnished equipment. It NEVER converts visual transcripts,
archived procurement or an apparent winner into signed supply revenue.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path


def _file_sha(path: Path, max_size: int = 25_000_000) -> str:
    if (path.is_symlink() or not path.is_file()
            or not 0 < path.stat().st_size <= max_size):
        raise ValueError("Missing, oversized or unsafe original platform file")
    digest = hashlib.sha256()
    with path.open("rb") as inp:
        for chunk in iter(lambda: inp.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit(*, runtime: Path, review: Path, max_cases: int = 20) -> dict:
    if runtime.is_symlink() or review.is_symlink() or not 1 <= max_cases <= 100:
        raise ValueError("Unsafe native commercial ETP scope")
    data = json.loads(review.read_text(encoding="utf-8"))
    if (data.get("schema_version") != "pq-first-party-etp-review/1"
            or not isinstance(data.get("cases"), list)
            or len(data["cases"]) > max_cases):
        raise ValueError("Unexpected source-verified ETP review schema")
    results = []
    seen = set()
    for entry in data["cases"]:
        number = entry.get("id")
        if not isinstance(number, str) or number in seen:
            raise ValueError("Duplicate or unsafe original commercial identifier")
        seen.add(number)
        if re.fullmatch(r"COM[0-9]{11}", number):
            card = runtime / "roseltorg-original-pages" / number / "original.html"
            card_sha = _file_sha(card, 900_000)
            if card_sha != entry.get("original_card_sha"):
                raise ValueError("Roseltorg original card source SHA mismatch")
            html = card.read_text(encoding="utf-8", errors="replace")
            if number not in html or "com.roseltorg.ru/file/get" not in html:
                raise ValueError("Missing original Roseltorg identity/source links")
            source_dir = runtime / "roseltorg-original-files" / number
            expected = entry.get("original_cancel_pdf_sha") or entry.get("original_results_pdf_sha")
            candidates = [
                p for p in source_dir.glob("*.pdf")
                if re.fullmatch(r"[0-9]{4,12}\.pdf", p.name)
                and _file_sha(p) == expected
                and f"/id/{p.stem}/" in html
            ]
            if len(candidates) != 1 or candidates[0].read_bytes()[:4] != b"%PDF":
                raise ValueError("Original Roseltorg source PDF not bound to card")
            outcome = "visual_review_provisional"
            if entry.get("signed_contract") is not False:
                raise ValueError("Roseltorg PDF cannot certify signed contract")
            findings = {
                "source_pdf_sha256": expected,
                "source_pdf_member_id": candidates[0].stem,
                "visual_review_scope": entry.get("class"),
                "source_outcome_scope": outcome,
                "equipment_supply_contract_verified": False,
                "voltage_6_35_verified_from_native_text": False,
            }
        elif re.fullmatch(r"RAD[0-9]{9}", number):
            card = runtime / "rad-originals" / number / "card.html"
            card_sha = _file_sha(card, 900_000)
            if card_sha != entry.get("original_card_sha"):
                raise ValueError("RAD original card source SHA mismatch")
            html = card.read_text(encoding="utf-8", errors="replace")
            if number not in html or "/etp/downloadppf?uuid=" not in html:
                raise ValueError("Missing original RAD procedure/files link")
            original = runtime / "rad-originals" / number / "technical.docx"
            sha = _file_sha(original, 3_000_000)
            if sha != entry.get("original_technical_docx_sha"):
                raise ValueError("RAD native technical DOCX SHA mismatch")
            from docx import Document

            document = Document(original)
            text = "\n".join(
                [p.text for p in document.paragraphs]
                + [" | ".join(c.text for c in row.cells)
                   for table in document.tables for row in table.rows])
            if not all(
                " ".join(str(quote).split()) in " ".join(text.split())
                for quote in entry.get("native_quotes", [])
            ):
                raise ValueError("RAD 35kV technical original quote not found in DOCX")
            if not entry.get("native_quotes") or entry.get("owner_furnished") is not True:
                # The outside-v29 parent repeat may have different facts, but
                # a missing reviewed owner-furnished quote cannot certify supply.
                if entry.get("same_project_as"):
                    findings = {
                        "source_docx_sha256": sha,
                        "source_outcome_scope": "parent_project_duplicate_association_only",
                        "equipment_supply_contract_verified": False,
                        "voltage_6_35_verified_from_native_text": False,
                    }
                else:
                    raise ValueError("RAD 35kV equipment ownership evidence missing")
            else:
                findings = {
                    "source_docx_sha256": sha,
                    "source_outcome_scope": "35kv_design_install_owner_furnished_bsk",
                    "equipment_supply_contract_verified": False,
                    "voltage_6_35_verified_from_native_text": True,
                }
            if entry.get("signed_contract") is not False:
                raise ValueError("RAD technical DOCX cannot certify signed contract")
        else:
            raise ValueError("Unknown native original commercial procedure type")
        results.append({
            "procedure_id": number,
            "source_card_sha256": card_sha,
            "source_level": "first_party_platform_page_and_original_document_sha_verified",
            "signed_contract_verified": False,
            "signed_supplier_verified": False,
            **findings,
        })
    return {
        "schema_version": "pq-original-commercial-overlay-audit/1",
        "source_case_count": len(results),
        "signed_suppliers_verified": 0,
        "equipment_only_sales_verified": 0,
        "records": results,
    }
