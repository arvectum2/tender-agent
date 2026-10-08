from __future__ import annotations

import hashlib
import os
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.modules.application_package_generator.models import (
    ApplicationDraftFieldProvenance,
    ApplicationDraftGeneration,
)
from src.modules.application_package_generator.schemas import (
    ApplicationDraftFieldMapping,
    GenerateApplicationDraftRequest,
)
from src.modules.document_store.models import ArtifactVersion, DocumentArtifact
from src.modules.document_store.schemas import (
    AddArtifactVersionRequest,
    CreateArtifactRequest,
    LinkArtifactRequest,
)
from src.modules.document_store.service import (
    add_artifact_version,
    create_artifact,
    get_artifact,
    link_artifact,
)
from src.shared.enums import ArtifactType
from src.shared.errors import NotFoundError, ValidationError
from src.tender_research.config import load_config

_DOCX = ".docx"
_XLSX = ".xlsx"
_DOCX_PREFIX = "docx:token:"
_XLSX_PREFIX = "xlsx:cell:"
_REVIEW_STATE = "DRAFT_REVIEW_ONLY"
_SAFE_SEGMENT = re.compile(r"[^A-Za-z0-9_.-]+")


@dataclass(frozen=True)
class RenderedField:
    source_type: str
    source_ref: str
    target_locator: str
    original_value: object | None
    generated_value: object | None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _data_root() -> Path:
    root = Path(load_config().data_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _path_under_data_root(storage_uri: str) -> Path:
    root = _data_root()
    raw = str(storage_uri or "").strip()
    if raw.startswith("file://"):
        candidate = Path(unquote(raw[7:]))
    elif "://" in raw:
        raise ValidationError(
            "Template artifact must use an internal local file storage URI"
        )
    else:
        candidate = Path(raw)
        if not candidate.is_absolute():
            candidate = root / candidate
    resolved = candidate.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise ValidationError(
            "Template artifact path must stay under the configured data directory"
        ) from exc
    if not resolved.is_file() or resolved.is_symlink():
        raise ValidationError("Template artifact file is missing or unsafe")
    return resolved


def _safe_segment(value: str) -> str:
    normalized = _SAFE_SEGMENT.sub("-", str(value or "").strip()).strip(".-")
    return normalized[:96] or "draft"


def _output_file_name(role: str, suffix: str) -> str:
    if role == "COMMERCIAL_PROPOSAL" and suffix == _XLSX:
        return "working_kp.xlsx"
    names = {
        "FORM_2": "form_2",
        "COMMERCIAL_PROPOSAL": "commercial_proposal",
        "DECLARATION": "declaration",
        "SPECIFICATION": "specification",
        "INVENTORY": "inventory",
    }
    return f"{names[role]}{suffix}"


def _lineage_key(deal_id: str, role: str, suffix: str) -> str:
    if role == "COMMERCIAL_PROPOSAL" and suffix == _XLSX:
        return f"{deal_id}:working_kp.xlsx"
    return f"{deal_id}:{role.lower()}:{suffix.lstrip('.')}"


def _template_version(
    session: Session,
    *,
    artifact: DocumentArtifact,
    expected_version: int,
) -> ArtifactVersion:
    if artifact.current_version != expected_version:
        raise ValidationError(
            f"Template artifact version changed: expected {expected_version}, current {artifact.current_version}"
        )
    version = session.scalar(
        select(ArtifactVersion).where(
            ArtifactVersion.artifact_ref == artifact.artifact_ref,
            ArtifactVersion.version_no == expected_version,
        )
    )
    if not version:
        raise ValidationError("Template artifact version evidence is missing")
    if version.storage_uri != artifact.storage_uri:
        raise ValidationError(
            "Template artifact/version storage lineage is inconsistent"
        )
    return version


def _validate_template(
    session: Session,
    payload: GenerateApplicationDraftRequest,
) -> tuple[DocumentArtifact, Path, str, str]:
    artifact = get_artifact(session, payload.template_artifact_ref)
    if artifact.deal_id not in {None, payload.deal_id}:
        raise ValidationError("Template artifact belongs to another deal")
    version = _template_version(
        session,
        artifact=artifact,
        expected_version=payload.expected_template_version,
    )
    path = _path_under_data_root(version.storage_uri)
    suffix = path.suffix.lower()
    if suffix not in {_DOCX, _XLSX}:
        raise ValidationError("Only DOCX and XLSX client templates are supported")
    if suffix == _DOCX and any(
        not item.target_locator.startswith(_DOCX_PREFIX) for item in payload.fields
    ):
        raise ValidationError(
            "DOCX templates require docx:token:<literal> target locators"
        )
    if suffix == _XLSX and any(
        not item.target_locator.startswith(_XLSX_PREFIX) for item in payload.fields
    ):
        raise ValidationError(
            "XLSX templates require xlsx:cell:<sheet>!<cell> target locators"
        )
    actual_sha = _sha256(path)
    expected_sha = version.checksum_sha256 or artifact.checksum_sha256
    if expected_sha and expected_sha.lower() != actual_sha:
        raise ValidationError(
            "Template artifact checksum no longer matches its persisted version"
        )
    return artifact, path, suffix, actual_sha


def _validate_source(
    session: Session,
    *,
    deal_id: str,
    mapping: ApplicationDraftFieldMapping,
) -> tuple[str, str]:
    source_type = mapping.source_type.strip().upper()
    source_ref = mapping.source_ref.strip()
    if source_type == "DEAL" and source_ref != deal_id:
        raise ValidationError("DEAL field provenance must reference the current deal")
    if source_type == "ARTIFACT":
        artifact = get_artifact(session, source_ref)
        if artifact.deal_id not in {None, deal_id}:
            raise ValidationError("ARTIFACT field provenance belongs to another deal")
    return source_type, source_ref


def _iter_docx_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested_table in cell.tables:
                    for nested_row in nested_table.rows:
                        for nested_cell in nested_row.cells:
                            yield from nested_cell.paragraphs


def _render_docx(
    template_path: Path,
    output_path: Path,
    fields: list[tuple[ApplicationDraftFieldMapping, str, str]],
) -> list[RenderedField]:
    document = Document(template_path)
    rendered: list[RenderedField] = []
    paragraphs = list(_iter_docx_paragraphs(document))
    for mapping, source_type, source_ref in fields:
        token = mapping.target_locator[len(_DOCX_PREFIX) :]
        if not token:
            raise ValidationError("DOCX token target must not be empty")
        matches = []
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if token in run.text:
                    matches.append(run)
        total_occurrences = sum(run.text.count(token) for run in matches)
        if total_occurrences != 1:
            raise ValidationError(
                f"DOCX target '{mapping.target_locator}' must resolve to exactly one in-run token occurrence"
            )
        run = matches[0]
        original = token
        generated = "" if mapping.value is None else str(mapping.value)
        run.text = run.text.replace(token, generated, 1)
        rendered.append(
            RenderedField(
                source_type=source_type,
                source_ref=source_ref,
                target_locator=mapping.target_locator,
                original_value=original,
                generated_value=mapping.value,
            )
        )
    document.save(output_path)
    return rendered


def _parse_xlsx_locator(locator: str) -> tuple[str, str]:
    payload = locator[len(_XLSX_PREFIX) :]
    if "!" not in payload:
        raise ValidationError("XLSX target must be xlsx:cell:<sheet>!<cell>")
    sheet, coordinate = payload.rsplit("!", 1)
    sheet = sheet.strip()
    coordinate = coordinate.strip().upper()
    if not sheet or not coordinate:
        raise ValidationError("XLSX sheet and cell must not be empty")
    try:
        coordinate_from_string(coordinate)
    except ValueError as exc:
        raise ValidationError(f"Invalid XLSX cell coordinate '{coordinate}'") from exc
    return sheet, coordinate


def _merged_anchor_or_fail(worksheet, coordinate: str) -> None:
    for merged in worksheet.merged_cells.ranges:
        if coordinate in merged:
            if coordinate != merged.start_cell.coordinate:
                raise ValidationError(
                    f"XLSX target '{worksheet.title}!{coordinate}' is not the anchor of merged range {merged}"
                )
            return


def _render_xlsx(
    template_path: Path,
    output_path: Path,
    fields: list[tuple[ApplicationDraftFieldMapping, str, str]],
) -> list[RenderedField]:
    workbook = load_workbook(template_path, data_only=False, keep_links=True)
    rendered: list[RenderedField] = []
    for mapping, source_type, source_ref in fields:
        sheet_name, coordinate = _parse_xlsx_locator(mapping.target_locator)
        if sheet_name not in workbook.sheetnames:
            raise ValidationError(f"XLSX target sheet '{sheet_name}' does not exist")
        worksheet = workbook[sheet_name]
        _merged_anchor_or_fail(worksheet, coordinate)
        cell = worksheet[coordinate]
        original = cell.value
        cell.value = mapping.value
        rendered.append(
            RenderedField(
                source_type=source_type,
                source_ref=source_ref,
                target_locator=mapping.target_locator,
                original_value=original,
                generated_value=mapping.value,
            )
        )
    workbook.save(output_path)
    return rendered


def _latest_generation(
    session: Session,
    *,
    deal_id: str,
    lineage_key: str,
) -> ApplicationDraftGeneration | None:
    return session.scalar(
        select(ApplicationDraftGeneration)
        .where(
            ApplicationDraftGeneration.deal_id == deal_id,
            ApplicationDraftGeneration.lineage_key == lineage_key,
        )
        .order_by(ApplicationDraftGeneration.generation_version_no.desc())
        .limit(1)
    )


def _provenance_rows(
    session: Session,
    generation_id: str,
) -> list[ApplicationDraftFieldProvenance]:
    return list(
        session.scalars(
            select(ApplicationDraftFieldProvenance)
            .where(ApplicationDraftFieldProvenance.generation_id == generation_id)
            .order_by(
                ApplicationDraftFieldProvenance.created_at.asc(),
                ApplicationDraftFieldProvenance.id.asc(),
            )
        )
    )


def generate_application_draft(
    session: Session,
    payload: GenerateApplicationDraftRequest,
) -> tuple[
    ApplicationDraftGeneration,
    DocumentArtifact,
    list[ApplicationDraftFieldProvenance],
]:
    template, template_path, suffix, template_sha = _validate_template(session, payload)
    validated_fields = [
        (mapping, *_validate_source(session, deal_id=payload.deal_id, mapping=mapping))
        for mapping in payload.fields
    ]
    lineage_key = _lineage_key(payload.deal_id, payload.document_role, suffix)
    previous = _latest_generation(
        session, deal_id=payload.deal_id, lineage_key=lineage_key
    )
    if previous:
        output_artifact = get_artifact(session, previous.output_artifact_ref)
        if output_artifact.deal_id != payload.deal_id:
            raise ValidationError("Generated artifact lineage belongs to another deal")
        version_no = output_artifact.current_version + 1
    else:
        output_artifact = None
        version_no = 1

    generation_id = str(uuid4())
    file_name = _output_file_name(payload.document_role, suffix)
    relative_dir = (
        Path("application-drafts")
        / _safe_segment(payload.deal_id)
        / _safe_segment(lineage_key)
        / f"v{version_no:04d}"
    )
    output_dir = _data_root() / relative_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / file_name
    partial_path = output_path.with_suffix(output_path.suffix + ".partial")
    partial_path.unlink(missing_ok=True)

    try:
        if suffix == _DOCX:
            rendered = _render_docx(template_path, partial_path, validated_fields)
            mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            output_format = "docx"
        else:
            rendered = _render_xlsx(template_path, partial_path, validated_fields)
            mime_type = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            output_format = "xlsx"
        with partial_path.open("rb") as handle:
            os.fsync(handle.fileno())
        os.replace(partial_path, output_path)
    finally:
        partial_path.unlink(missing_ok=True)

    output_sha = _sha256(output_path)
    storage_uri = output_path.relative_to(_data_root()).as_posix()
    if output_artifact is None:
        output_artifact = create_artifact(
            session,
            CreateArtifactRequest(
                deal_id=payload.deal_id,
                artifact_type=ArtifactType.GENERATED_DOC,
                file_name=file_name,
                mime_type=mime_type,
                storage_uri=storage_uri,
                checksum_sha256=output_sha,
            ),
        )
        if output_artifact.current_version != version_no:
            raise ValidationError("Generated artifact version lineage is inconsistent")
    else:
        if output_artifact.file_name != file_name:
            raise ValidationError(
                "Generated artifact canonical file name changed inside one lineage"
            )
        artifact_version = add_artifact_version(
            session,
            output_artifact.artifact_ref,
            AddArtifactVersionRequest(
                storage_uri=storage_uri,
                checksum_sha256=output_sha,
            ),
        )
        if artifact_version.version_no != version_no:
            raise ValidationError("Generated artifact version lineage is inconsistent")
        output_artifact = get_artifact(session, output_artifact.artifact_ref)

    generation = ApplicationDraftGeneration(
        id=generation_id,
        deal_id=payload.deal_id,
        template_artifact_ref=template.artifact_ref,
        output_artifact_ref=output_artifact.artifact_ref,
        document_role=payload.document_role,
        output_format=output_format,
        lineage_key=lineage_key,
        template_version_no=payload.expected_template_version,
        generation_version_no=version_no,
        template_sha256=template_sha,
        output_sha256=output_sha,
        output_file_name=file_name,
        output_storage_uri=storage_uri,
        review_state=_REVIEW_STATE,
        signature_performed=False,
        submission_performed=False,
        external_delivery_performed=False,
    )
    session.add(generation)
    for field in rendered:
        session.add(
            ApplicationDraftFieldProvenance(
                generation_id=generation_id,
                source_type=field.source_type,
                source_ref=field.source_ref,
                target_locator=field.target_locator,
                original_value_json=field.original_value,
                generated_value_json=field.generated_value,
            )
        )
    session.flush()
    link_artifact(
        session,
        output_artifact.artifact_ref,
        LinkArtifactRequest(
            linked_object_type="APPLICATION_DRAFT_GENERATION",
            linked_object_ref=generation_id,
        ),
    )
    session.refresh(generation)
    return generation, output_artifact, _provenance_rows(session, generation.id)


def get_application_draft(
    session: Session,
    generation_id: str,
) -> tuple[
    ApplicationDraftGeneration,
    DocumentArtifact,
    list[ApplicationDraftFieldProvenance],
]:
    generation = session.get(ApplicationDraftGeneration, generation_id)
    if not generation:
        raise NotFoundError(
            f"Application draft generation '{generation_id}' was not found"
        )
    artifact = get_artifact(session, generation.output_artifact_ref)
    return generation, artifact, _provenance_rows(session, generation.id)


def list_application_drafts(
    session: Session,
    *,
    deal_id: str | None = None,
    document_role: str | None = None,
) -> list[
    tuple[
        ApplicationDraftGeneration,
        DocumentArtifact,
        list[ApplicationDraftFieldProvenance],
    ]
]:
    query = select(ApplicationDraftGeneration).order_by(
        ApplicationDraftGeneration.created_at.desc(),
        ApplicationDraftGeneration.id.desc(),
    )
    if deal_id:
        query = query.where(ApplicationDraftGeneration.deal_id == deal_id)
    if document_role:
        query = query.where(ApplicationDraftGeneration.document_role == document_role)
    generations = list(session.scalars(query))
    return [
        (
            generation,
            get_artifact(session, generation.output_artifact_ref),
            _provenance_rows(session, generation.id),
        )
        for generation in generations
    ]
