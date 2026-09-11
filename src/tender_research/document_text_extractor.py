from __future__ import annotations

import io
import os
import subprocess
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree


EXTRACTED_STATUS = "extracted"
FAILED_STATUS = "failed"
UNSUPPORTED_STATUS = "unsupported"
EMPTY_STATUS = "empty"
_MACOS_TEXTUTIL = Path("/usr/bin/textutil")
_TEXTUTIL_TIMEOUT_SECONDS = 30
# Bounded, fixed locations for the wvHtml legacy-Word converter.  The binary
# is never resolved through PATH and never influenced by document content.
_WVHTML_CANDIDATES = (
    Path("/opt/homebrew/bin/wvHtml"),
    Path("/usr/local/bin/wvHtml"),
    Path("/usr/bin/wvHtml"),
)
_WVHTML_TIMEOUT_SECONDS = 30
# Generated HTML is parsed, never rendered.  The bound guards the HTML parser
# against pathological converter output; ordinary procurement documents are
# orders of magnitude smaller.
_WVHTML_MAX_HTML_BYTES = 16 * 1024 * 1024
_SUPPORTED_EXTENSIONS = (
    ".txt",
    ".doc",
    ".docx",
    ".rtf",
    ".pdf",
    ".xlsx",
    ".xls",
    ".html",
    ".htm",
    ".xml",
    ".csv",
    ".json",
)


def extract_text(local_path: str, max_chars: int = 2_000_000) -> tuple[str, str]:
    declared_ext = Path(local_path).suffix.lower()
    try:
        with open(local_path, "rb") as f:
            content = f.read()
    except OSError as e:
        return FAILED_STATUS, f"File read error: {e}"
    if not content:
        return EMPTY_STATUS, ""

    # Public EIS attachments are not guaranteed to have a trustworthy suffix.
    # Prefer deterministic file signatures/container members where they identify
    # a supported format unambiguously; otherwise preserve the declared suffix.
    detected_ext = _sniff_supported_extension(content)
    ext = detected_ext or declared_ext
    result = _extract_by_ext(ext, content, max_chars, local_path=local_path)
    if result is None or not result.strip():
        if _is_unsupported_ext(ext):
            return UNSUPPORTED_STATUS, (result or "")
        return EMPTY_STATUS, (result or "")
    return EXTRACTED_STATUS, result[:max_chars]


def _sniff_supported_extension(content: bytes) -> str | None:
    """Identify supported PDF/OOXML formats without trusting the filename.

    Only deterministic signatures are accepted.  ZIP containers are inspected
    through their canonical OOXML member names and are never unpacked here.
    """

    if content.startswith(b"%PDF-"):
        return ".pdf"
    if not content.startswith(b"PK"):
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            try:
                archive.getinfo("word/document.xml")
                return ".docx"
            except KeyError:
                pass
            try:
                archive.getinfo("xl/workbook.xml")
                return ".xlsx"
            except KeyError:
                pass
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile):
        return None
    return None


def _extract_by_ext(
    ext: str,
    content: bytes,
    max_chars: int,
    *,
    local_path: str | None = None,
) -> str | None:
    if ext == ".txt":
        return _extract_txt(content)
    if ext in (".doc", ".rtf"):
        return (
            _extract_legacy_office(local_path, max_chars, ext=ext)
            if local_path is not None
            else ""
        )
    if ext == ".docx":
        return _extract_docx(content)
    if ext == ".pdf":
        return _extract_pdf(content, max_chars)
    if ext in (".xlsx", ".xls"):
        return _extract_xlsx(content)
    if ext in (".html", ".htm"):
        return _extract_html(content)
    if ext == ".xml":
        return _extract_xml(content)
    if ext == ".csv":
        return _extract_txt(content)
    if ext == ".json":
        return _extract_txt(content)
    return None


def _is_unsupported_ext(ext: str) -> bool:
    return ext not in _SUPPORTED_EXTENSIONS


def _extract_txt(content: bytes) -> str:
    for enc in ("utf-8", "cp1251", "koi8-r", "latin-1"):
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="replace")


def _extract_legacy_office(local_path: str, max_chars: int, *, ext: str = ".doc") -> str:
    """Extract legacy Word text, preferring table structure for .doc files.

    Legacy ``.doc`` files first attempt a structure-preserving wvHtml
    conversion whose tables are projected into the same tab-separated row
    convention used by native DOCX/XLSX extraction.  Any wvHtml failure
    (missing binary, timeout, nonzero exit, oversized or table-less output)
    falls back to the pre-existing textutil plain-text conversion.  ``.rtf``
    keeps the textutil path.  No table geometry is ever inferred from flat
    prose: only real ``<tr>`` rows with real ``<td>``/``<th>`` cells qualify.
    """

    if ext == ".doc":
        structured = _extract_legacy_word_with_wvhtml(local_path, max_chars)
        if structured:
            return structured
    if not _MACOS_TEXTUTIL.is_file():
        return ""
    try:
        completed = subprocess.run(
            [
                str(_MACOS_TEXTUTIL),
                "-convert",
                "txt",
                "-stdout",
                "-encoding",
                "UTF-8",
                local_path,
            ],
            capture_output=True,
            check=False,
            timeout=_TEXTUTIL_TIMEOUT_SECONDS,
        )
    except (OSError, subprocess.SubprocessError):
        return ""
    if completed.returncode != 0 or not completed.stdout:
        return ""
    return _extract_txt(completed.stdout)[:max_chars]


def _resolve_wvhtml() -> Path | None:
    """Return the first usable fixed-location wvHtml binary, if any."""

    for candidate in _WVHTML_CANDIDATES:
        if candidate.is_file():
            return candidate
    return None


def _extract_legacy_word_with_wvhtml(local_path: str, max_chars: int) -> str:
    """Project legacy Word tables through wvHtml into tab-separated rows.

    The source file is copied into a temporary directory first because wvHtml
    may emit sidecar files (e.g. ``.wmf``) next to its inputs; the original
    stays read-only and every generated file is discarded with the directory.
    Returns an empty string unless at least one multi-cell table row survives
    projection, in which case callers fall back to plain-text extraction.
    """

    executable = _resolve_wvhtml()
    if executable is None:
        return ""
    try:
        with open(local_path, "rb") as handle:
            source_bytes = handle.read()
    except OSError:
        return ""
    if not source_bytes:
        return ""
    try:
        with tempfile.TemporaryDirectory(prefix="wvhtml-") as workdir:
            work_path = Path(workdir)
            input_path = work_path / f"source{Path(local_path).suffix.lower() or '.doc'}"
            output_path = work_path / "converted.html"
            try:
                input_path.write_bytes(source_bytes)
            except OSError:
                return ""
            try:
                completed = subprocess.run(
                    [str(executable), str(input_path), str(output_path)],
                    capture_output=True,
                    check=False,
                    timeout=_WVHTML_TIMEOUT_SECONDS,
                )
            except (OSError, subprocess.SubprocessError):
                return ""
            # A nonzero exit (wvHtml is known to segfault on some files) may
            # still leave partial HTML behind.  Never consume it.
            if completed.returncode != 0:
                return ""
            try:
                if not output_path.is_file():
                    return ""
                if output_path.stat().st_size == 0:
                    return ""
                if output_path.stat().st_size > _WVHTML_MAX_HTML_BYTES:
                    return ""
                html_bytes = output_path.read_bytes()
            except OSError:
                return ""
            return _project_wvhtml_tables(html_bytes, max_chars)
    except OSError:
        return ""


class _WvHtmlTableProjector(HTMLParser):
    """Project wvHtml tables into the shared tab-separated row convention.

    Table rows become ``cell1<TAB>cell2...`` lines; prose paragraphs outside
    tables remain newline-separated text.  Script/style content is ignored,
    entities are decoded by the parser, and nested inline markup inside a
    cell contributes its text.  Links, images and converter sidecars are
    never interpreted.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.table_row_count = 0
        self._table_depth = 0
        self._current_row: list[str] | None = None
        self._current_cell: list[str] | None = None
        self._current_paragraph: list[str] | None = None
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        normalized = tag.lower()
        if normalized in ("script", "style"):
            self._ignored_depth += 1
        elif normalized == "table":
            self._table_depth += 1
        elif normalized == "tr" and self._table_depth:
            self._current_row = []
        elif normalized in ("td", "th") and self._current_row is not None:
            self._current_cell = []
        elif normalized == "p" and not self._table_depth:
            self._current_paragraph = []
        elif normalized == "br":
            self._flush_paragraph()

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.lower()
        if normalized in ("script", "style"):
            self._ignored_depth = max(0, self._ignored_depth - 1)
        elif normalized in ("td", "th") and self._current_cell is not None:
            self._current_row.append(" ".join("".join(self._current_cell).split()))
            self._current_cell = None
        elif normalized == "tr" and self._current_row is not None:
            cells = [cell for cell in self._current_row if cell]
            if len(cells) >= 2:
                self.lines.append("\t".join(cells))
                self.table_row_count += 1
            elif cells:
                # A single-cell table row carries no column structure; keep
                # its text as ordinary prose instead of dropping it.
                self.lines.append(cells[0])
            self._current_row = None
        elif normalized == "table":
            self._table_depth = max(0, self._table_depth - 1)
        elif normalized == "p":
            self._flush_paragraph()

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        if self._current_cell is not None:
            self._current_cell.append(data)
        elif self._current_paragraph is not None:
            self._current_paragraph.append(data)

    def _flush_paragraph(self) -> None:
        if self._current_paragraph is not None:
            text = " ".join("".join(self._current_paragraph).split())
            if text:
                self.lines.append(text)
            self._current_paragraph = None


def _project_wvhtml_tables(html_bytes: bytes, max_chars: int) -> str:
    """Return projected text, or "" when no multi-cell table row survives."""

    try:
        html = html_bytes.decode("utf-8", errors="replace")
    except (LookupError, ValueError):
        return ""
    projector = _WvHtmlTableProjector()
    try:
        projector.feed(html)
    except Exception:  # noqa: BLE001 - any malformed converter HTML fails closed
        return ""
    projector._flush_paragraph()
    if projector.table_row_count < 1:
        return ""
    return "\n".join(projector.lines)[:max_chars]


def _extract_docx(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml_content = z.read("word/document.xml")
        root = ElementTree.fromstring(xml_content)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        body = root.find("w:body", ns)
        if body is None:
            return ""

        blocks: list[str] = []
        for child in body:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph_text = "".join(
                    t.text or ""
                    for t in child.iter(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                    )
                ).strip()
                if paragraph_text:
                    blocks.append(paragraph_text)
            elif tag == "tbl":
                for row in child.findall("w:tr", ns):
                    cells: list[str] = []
                    for cell in row.findall("w:tc", ns):
                        cell_text = "".join(
                            t.text or ""
                            for t in cell.iter(
                                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                            )
                        ).strip()
                        cells.append(cell_text)
                    normalized_cells = [cell for cell in cells if cell]
                    if normalized_cells:
                        blocks.append("\t".join(normalized_cells))
        return "\n".join(blocks)
    except Exception:
        return ""


def _extract_pdf(content: bytes, max_chars: int) -> str:
    try:
        import pypdf
    except ImportError:
        return ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        texts = []
        for i, page in enumerate(reader.pages):
            if i >= 10:
                break
            t = page.extract_text() or ""
            texts.append(t)
        return " ".join(texts)[:max_chars]
    except Exception:
        return ""


def _extract_xlsx(content: bytes) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_lines = []
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) if v is not None else "" for v in row]
                sheet_lines.append("\t".join(vals))
            if sheet_lines:
                lines.append(f"=== {sheet_name} ===")
                lines.extend(sheet_lines)
        wb.close()
        return "\n".join(lines)
    except Exception:
        return ""


def _extract_html(content: bytes) -> str:
    text = _extract_txt(content)
    import re

    clean = re.sub(
        r"<script[^>]*>.*?</script>",
        "",
        text,
        flags=re.DOTALL | re.IGNORECASE,
    )
    clean = re.sub(
        r"<style[^>]*>.*?</style>",
        "",
        clean,
        flags=re.DOTALL | re.IGNORECASE,
    )
    clean = re.sub(r"<[^>]+>", " ", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    return clean[:500_000]


def _extract_xml(content: bytes) -> str:
    """Return deterministic, parseable XML rather than destroying element names.

    Downstream EIS metadata and purchase-object extractors are intentionally
    namespace-agnostic but require the XML element structure.  The previous
    plain-text projection stripped every tag, making authoritative EIS fields
    such as ``maxPrice``, ``purchaseObject`` and ``OKPD2`` impossible to recover.
    Parsing and re-serializing also rejects malformed XML instead of passing an
    arbitrary tag-like string into structured extractors.
    """

    try:
        root = ElementTree.fromstring(content)
    except (ElementTree.ParseError, ValueError):
        try:
            text = _extract_txt(content)
            root = ElementTree.fromstring(text)
        except (ElementTree.ParseError, ValueError):
            return ""
    return ElementTree.tostring(root, encoding="unicode", method="xml")
