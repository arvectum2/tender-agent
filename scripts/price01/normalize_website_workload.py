"""PRICE-01 source-bound workload matrix: no inferred hours, price ratios or bids.

Original DOCX files stay on local ArvectumSSD. Export only attested short
clauses, SHA-256 and measurement states; unknowns are deliberately null.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from scripts.price01.audit_original_website_cohort import PINS, _normalized

_CASES = {
    "0158300034526000388": "AKSAY",
    "0119200000126013647": "KRASNOYARSK",
    "0377100005726000128": "CHEMAL",
}
_EXPECTED_SHA = {p.number: p.scope_sha256 for p in PINS if p.number in _CASES}
assert len(_EXPECTED_SHA) == 3


@dataclass(frozen=True)
class EvidenceRule:
    dimension: str
    source_phrase: str
    kind: str
    value: str | int
    unit: str
    interpretation: str


# Values here are *reviewed literal obligations*; a matching original clause
# and its exact file hash are mandatory to attest each. No economic conversion.
_RULES: dict[str, tuple[EvidenceRule, ...]] = {
    "0158300034526000388": (
        EvidenceRule("delivery_working_days", "в течение 14 рабочих дней с момента заключения Контракта", "UP_TO", 14, "working_days", "Договорный срок выполнения; не означает количество человеко-дней."),
        EvidenceRule("cms_platform", "Конкретное наименование CMS Заказчиком не устанавливается", "SPECIFIED", "RUSSIAN_REGISTRY_CMS_CONTRACTOR_CHOICE", "category", "Российская CMS из реестра; конкретный продукт выбирает исполнитель."),
        EvidenceRule("design_basis", "основе предоставленного тёмного макета сайта Молодёжного центра", "SPECIFIED", "CLIENT_PROVIDED_MOCKUP_ADAPTED", "category", "Доработка по предоставленному макету; число итераций согласовывается отдельно."),
        EvidenceRule("initial_content_origin", "Заказчик предоставляет окончательные тексты, изображения", "SPECIFIED", "CUSTOMER_PROVIDED", "category", "Исходный контент предоставляет заказчик; точное количество страниц в ТЗ не установлено."),
        EvidenceRule("postacceptance_obligation", "дальнейшее сопровождение после приёмки в предмет работ не входят", "SPECIFIED", "ROUTINE_SUPPORT_EXCLUDED", "category", "Регулярное сопровождение после приёмки исключено; это не доказательство отсутствия гарантий."),
        EvidenceRule("hosting_payment", "хостинг сроком на 12 месяцев, если они требуются", "CONDITIONAL", "CONTRACTOR_IF_REQUIRED_12_MONTHS", "category", "Лицензию CMS и хостинг на год при необходимости оплачивает исполнитель; стоимость не указана."),
        EvidenceRule("knowledge_transfer", "не менее 2 часов с ответами на вопросы", "AT_LEAST", 2, "hours", "Одна демонстрация и передача знаний продолжительностью не менее 2 часов."),
        EvidenceRule("site_purpose", "информирования посетителей о новостях, мероприятиях, пространствах и проектах", "SPECIFIED", "CONTENT_AND_EVENTS_SITE", "category", "Сайт с новостями, событиями, каталогом проектов; чужие поисковые модули не приравниваются."),
    ),
    "0119200000126013647": (
        EvidenceRule("delivery_working_days", "в течение 10 (десяти) рабочих дней с даты заключения контракта", "UP_TO", 10, "working_days", "Срок договора; не определяет количество человеко-дней."),
        EvidenceRule("cms_platform", "установка и настройка отечественной CMS", "SPECIFIED", "RUSSIAN_CMS_PRODUCT_NOT_PINNED", "category", "Отечественная CMS обязательна, но точная лицензия и её цена неизвестны."),
        EvidenceRule("design_basis", "разработка интерактивного кликабельного прототипа всех страниц (Figma)", "SPECIFIED", "NEW_FIGMA_CLICKABLE_ALL_PAGES", "category", "Новый интерактивный Figma-прототип всех страниц и состояний ошибок, не готовый макет заказчика."),
        EvidenceRule("initial_content_origin", "загрузка первичных текстов от Заказчика", "SPECIFIED", "CUSTOMER_PROVIDED_PLUS_NEWS_AUTHORING", "category", "Контент заказчика и дополнительная обязанность создавать реальные новости."),
        EvidenceRule("postacceptance_obligation", "Срок сопровождения: 12 месяцев с момента подписания", "SPECIFIED", "ONGOING_SUPPORT_12_MONTHS_WITH_SLA", "category", "12 месяцев технического сопровождения с SLA и мелкими правками; не только гарантия дефектов."),
        EvidenceRule("support_critical_reaction_hours", "Критический | Недоступность сайта, утечка данных | 1 час | 4 часа", "AT_MOST", 1, "hours", "SLA критических инцидентов: реакция 1 час, устранение 4 часа."),
        EvidenceRule("knowledge_transfer", "семинара (до 2 часов) для сотрудников Заказчика", "UP_TO", 2, "hours", "Обучающий семинар продолжительностью до 2 часов."),
        EvidenceRule("site_purpose", "Строка поиска в шапке или над основным контентом", "SPECIFIED", "LEGAL_HELP_SEARCH_FORMS", "category", "Поиск по правовым материалам и новостям, отдельные формы обратной связи."),
        EvidenceRule("feedback_form", "поля: ФИО (Обязательное), Электронная почта (Обязательное)", "SPECIFIED", "PERSONAL_DATA_CONTACT_FORM", "category", "Форма ФИО/email с согласием ПДн и капчей; трудозатраты неизвестны."),
    ),
    "0377100005726000128": (
        EvidenceRule("delivery_working_days", "не более 21 рабочего дня с даты подписания контракта", "UP_TO", 21, "working_days", "Contractual deadline, not labor input."),
        EvidenceRule("cms_platform", "«1С-Битрикс: Управление сайтом» в редакции «Стандарт» сроком действия не менее 1 (одного) года", "SPECIFIED", "BITRIX_STANDARD_1_YEAR_PAID_BY_CONTRACTOR", "category", "Фиксирована лицензия «1С-Битрикс: Стандарт» минимум на год за счёт исполнителя; цена неизвестна."),
        EvidenceRule("design_basis", "сайт разрабатывается на базе готового специализированного шаблона", "SPECIFIED", "ONE_PREBUILT_TEMPLATE_CUSTOMIZATION", "category", "Настройка готового шаблона и одного варианта дизайна; не индивидуальный Figma-прототип."),
        EvidenceRule("initial_content_origin", "предоставленных Заказчиком", "SPECIFIED", "CUSTOMER_PROVIDED_UP_TO_40_PAGES", "category", "Материалы предоставляет заказчик; максимальный объём страниц зафиксирован отдельно."),
        EvidenceRule("initial_content_pages_max", "в объеме до 40 текстовых страниц", "UP_TO", 40, "text_pages", "Наполнение до 40 текстовых страниц заказчика, не 40 уникальных шаблонов."),
        EvidenceRule("postacceptance_obligation", "Гарантийный срок на разработанное программное обеспечение и выполненные работы составляет 12 (двенадцать) месяцев", "SPECIFIED", "DEFECT_WARRANTY_12_MONTHS", "category", "Гарантия на исправление дефектов 12 месяцев; не равнозначна годовому сопровождению."),
        EvidenceRule("hosting_payment", "услуги хостинга на первый год обслуживания в полном объеме входят в цену контракта", "SPECIFIED", "CONTRACTOR_HOSTING_AND_NEW_DOMAIN_FIRST_YEAR", "category", "Регистрация нового домена и первый год хостинга за счёт исполнителя; цены неизвестны."),
        EvidenceRule("site_purpose", "Карта объектов» реализуется Исполнителем на базе бесплатного API сервиса Яндекс.Карты", "SPECIFIED", "LOCATION_MAP_AND_SPEECH_ACCESSIBILITY", "category", "Интерактивная Яндекс.Карта объектов и функции доступности/озвучивания."),
        EvidenceRule("speech_accessibility", "через Web Speech API", "SPECIFIED", "WEB_SPEECH_API_AND_LOW_VISION_MODE", "category", "Озвучивание штатным Web Speech API и отдельные требования к версии для слабовидящих."),
    ),
}
assert {k: len(v) for k, v in _RULES.items()} == {k: (8 if k.endswith("388") else 9) for k in _RULES}


def extract_source_lines(docx_bytes: bytes) -> list[str]:
    if not 100 <= len(docx_bytes) <= 1_500_000 or not docx_bytes.startswith(b"PK"):
        raise ValueError("not a bounded Word source")
    try:
        document = Document(io.BytesIO(docx_bytes))
    except Exception as exc:
        raise ValueError("invalid source DOCX") from exc
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        lines.extend(
            " | ".join(cell.text.strip().replace("\n", "; ") for cell in row.cells)
            for row in table.rows
        )
    if not 10 <= len(lines) <= 2000:
        raise ValueError("unexpected original Word scope length")
    return lines


def attest_rule(lines: list[str], rule: EvidenceRule) -> dict:
    if not rule.source_phrase or not rule.dimension or not rule.interpretation:
        raise ValueError("missing reviewed requirement")
    matches = [
        (index, " ".join(text.split()))
        for index, text in enumerate(lines)
        if _normalized(rule.source_phrase) in _normalized(text)
    ]
    if len(matches) != 1:
        raise ValueError(f"scope dimension {rule.dimension} changed, missing or ambiguous")
    position, original = matches[0]
    return {
        "dimension": rule.dimension,
        "source_state": "ATTESTED_FROM_EXACT_DOCX",
        "constraint_kind": rule.kind,
        "source_value": rule.value,
        "unit": rule.unit,
        "source_quote": rule.source_phrase,
        "source_line_ordinal": position,
        "source_line_sha256": hashlib.sha256(original.encode("utf-8")).hexdigest(),
        "reviewed_interpretation": rule.interpretation,
    }


def audit_workload_case(number: str, data: bytes, *, expected_sha: str | None = None) -> dict:
    if number not in _CASES:
        raise ValueError("unknown procurement in workload comparison")
    sha = hashlib.sha256(data).hexdigest()
    if sha != (expected_sha or _EXPECTED_SHA[number]):
        raise ValueError("original DOCX SHA mismatch")
    lines = extract_source_lines(data)
    rows = [attest_rule(lines, rule) for rule in _RULES[number]]
    if len({r["dimension"] for r in rows}) != len(rows):
        raise ValueError("duplicate measured dimension")
    return {
        "registry_number": number,
        "case_key": _CASES[number],
        "original_technical_scope_docx_sha256": sha,
        "verified_dimensions": rows,
        "unknown_inputs": {
            "actual_labor_hours": None,
            "actual_labor_cost_rub": None,
            "cms_license_cost_rub": None,
            "hosting_and_domain_cost_rub": None,
            "optional_integrations_cost_rub": None,
            "actual_support_reserve_cost_rub": None,
            "num_custom_site_page_templates": None,
            "actual_initial_content_pages": None,
            "initial_content_pages_cap": 40 if number == "0377100005726000128" else None,
            "content_pages_scope": "UP_TO" if number == "0377100005726000128" else "NOT_QUANTIFIED_IN_REVIEWED_TZ",
        },
        "money_or_labor_conversion_estimated": False,
    }


def collect_workload(*, original_scope_root: Path) -> dict:
    cases = []
    for number in _CASES:
        file = original_scope_root / number / "technical-scope.docx"
        if not file.is_file() or file.is_symlink() or file.stat().st_size > 1_500_000:
            raise ValueError("canonical original DOCX missing or unsafe")
        cases.append(audit_workload_case(number, file.read_bytes()))
    deadlines = {row["registry_number"]: next(
        x["source_value"] for x in row["verified_dimensions"]
        if x["dimension"] == "delivery_working_days"
    ) for row in cases}
    if deadlines != {
        "0158300034526000388": 14, "0119200000126013647": 10,
        "0377100005726000128": 21,
    }:
        raise ValueError("contractual delivery times changed")
    return {
        "schema_version": "ta-price01-verified-workload-no-financial-coefficients-v1",
        "comparison_scope": "THREE_REVIEWED_ORIGINAL_TZ_WEBSITE_PROJECTS",
        "baseline_aksay_contracts_count": 1,
        "new_equivalent_website_contracts": 0,
        "other_related_site_contracts_requiring_cost_normalization": 2,
        "labor_hours_comparable": False,
        "scope_normalization_adjustment_rub": None,
        "cross_tender_market_price_rub": None,
        "bid_recommendation_rub": None,
        "company_costs_user_supplied": False,
        "monetary_multiplier_per_working_day": None,
        "cases": cases,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scope-root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    result = collect_workload(original_scope_root=args.scope_root)
    args.out.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print("WORKLOAD_CASES", len(result["cases"]), "COMPARABLE", result["baseline_aksay_contracts_count"])


if __name__ == "__main__":
    main()
