"""Immutable EIS/filestore source cohort audit; NEVER a mixed market bid forecast."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from statistics import median
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document


@dataclass(frozen=True)
class Pin:
    number: str
    zip_sha256: str
    final_sha256: str
    scope_sha256: str
    exact_scope_subject_fragment: str
    scope_evidence_phrases: tuple[str, ...]
    comparison: str
    reason: str


PINS = (
    Pin(
        "0158300034526000388",
        "d25187191599e705427676cc9de64ce665029381efa2490607d65ff237837bb8",
        "2b9f510c48e7d828c432e32b4dca3a4fff8edb75c66c3da37d46a12af3456af5",
        "0e40a162e85a69b496fb350e846a39c90695033fd5a46ebfb2683f9996132f7a",
        "Услуга по разработке веб сайта",
        ("в течение 14 рабочих дней", "российской CMS", "хостинг сроком на 12 месяцев"),
        "BASELINE_ONLY",
        "First primary website original, 45 admitted source-ranked prices. Keep one procurement as one basket.",
    ),
    Pin(
        "0132300001726000686",
        "9d7fb608f6f7d9853867bf84894ee88801bcac34316c6dd00f4f2047c31814c6",
        "9101c0ac598d949e44ef6db6b37fdf938dedf1eec04effb0d6c971879995087b",
        "e925f0c8e8543e073e8d9be0b0a0da2fc253edde274ef061a28d2b4b4ce8019c",
        "официального сайта городской Думы г. Дзержинск",
        ("интернет-приемной", "1 (одного) года", "лицензия ФСБ России"),
        "NOT_COMPARABLE_REGULATED_OFFICIAL_SITE",
        "Local authority e-reception and protected communications requirements plus annual support; only one admitted candidate.",
    ),
    Pin(
        "0119200000126013647",
        "03f28440a7d69cab318a768cddd8168481305eb87a6d884d042e6c3c2ac3ab9b",
        "58435ae386b0dda5d8b7f000b811e1c5ed09a3295cd2fcc41c81e93c44678d4e",
        "35110af58a89e8d1c77ac6077d675e146af826269f149463d7f0a4b730fdaac4",
        "корпоративного сайта",
        ("Figma", "отечественной CMS", "Поиск"),
        "RELATED_SCOPE_REQUIRES_NORMALIZATION",
        "Separate corporate site with original Figma prototype, search and legal bureau information model. Not one-for-one baseline.",
    ),
    Pin(
        "0377100005726000128",
        "8df59538a21209e6505961f09f8afd97646c6533aba67af05eb67bbc4dd4c69a",
        "99ae33db4fd603fcd440a4de9a4a7d3aeffbcda7d41112dd286ea10ce9e5c5ed",
        "a099f70cda44c5cdd52fe5ddcab6c462d24a5222f167d97d88d7e47b6a2481cd",
        'на базе "1С-Битрикс:Управление сайтом"',
        ("1С-Битрикс", "до 40 текстовых страниц", "Яндекс.Карты", "Web Speech API"),
        "RELATED_SCOPE_REQUIRES_NORMALIZATION",
        "Specified licensed Bitrix template, 40 pages, voice and mapping integration; different license/content basket.",
    ),
    Pin(
        "0372200203126000012",
        "6161b6fb9336398d0ba644af8c95e3236d887dbe08bd768caa95600e5bb14c20",
        "a9231c0125d31d27c8c22f5d4f2025fa8129a40bf1b60c3440ee41fb39b81b46",
        "3b2b8f96620e8a3aa90174b34ee47a272cb3506912796f9d0b8e2c45b1296664",
        "созданию сайта",
        ("CMS должна быть разработана специально для Заказчика", "PHP-фреймворка полного цикла", "Internet Explorer 11.0+"),
        "NOT_COMPARABLE_CUSTOM_CMS",
        "Custom proprietary-to-customer CMS in full-stack PHP, legacy-browser acceptance, 40 workdays; final EOK does not expose finalPrice.",
    ),
    Pin(
        "0884200000526000012",
        "dd9d3e108a3426e4f474d1ad2771c419f1354d85966839b5ad7f13b53768bafa",
        "b5d759b1b975b071ac94d2caec7137251051599b1507d8b615fc46dd8abbfae4",
        "fa3133414456b8413cfb2af1765e160603ee653ae874fd6354d708bc7208b815",
        "Создание и сопровождение сайта",
        ("1С-Битрикс", "двухфакторная аутентификация", "хранением журнала не менее 6 месяцев", "визуальный конструктор страниц"),
        "NOT_COMPARABLE_COMPOUND_SUPPORT",
        "Website plus continuing support, 2FA, six-month audit logging, advanced CMS builder; no like-for-like scope.",
    ),
)
assert len({p.number for p in PINS}) == len(PINS)


def _normalized(text: str) -> str:
    return " ".join(text.casefold().replace("ё", "е").split())


def _local(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _leaf(element: ET.Element, name: str) -> str | None:
    value = {(e.text or "").strip() for e in element.iter() if _local(e) == name and not len(e)}
    if len(value) != 1:
        return None
    return next(iter(value))


def _xml(raw: bytes) -> ET.Element:
    if not 1 <= len(raw) <= 2_000_000 or any(x in raw.upper() for x in (b"<!DOCTYPE", b"<!ENTITY")):
        raise ValueError("unsafe or oversized official XML")
    return ET.fromstring(raw)


def _notice_and_final(archive: ZipFile, number: str, final_sha: str) -> tuple[bytes, bytes, str, int]:
    notices = []
    finals = []
    for info in archive.infolist():
        if info.file_size > 2_000_000 or ".." in Path(info.filename).parts:
            raise ValueError("unsafe ZIP member")
        m = re.fullmatch(rf"epNotification[A-Za-z0-9_]+_{number}_(\d+)_([A-Za-z0-9]+)\.xml", info.filename)
        if m:
            notices.append((int(m.group(1)), info.filename))
        if re.fullmatch(rf"epProtocol[A-Za-z0-9]+Final_{number}_\d+_[A-Za-z0-9]+\.xml", info.filename):
            finals.append(info.filename)
    if not notices or len(finals) != 1:
        raise ValueError("missing or ambiguous original final/notice")
    latest = max(n for n, _ in notices)
    latest_names = [name for n, name in notices if n == latest]
    if len(latest_names) != 1:
        raise ValueError("ambiguous notice revision")
    notice_bytes = archive.read(latest_names[0])
    final_bytes = archive.read(finals[0])
    if hashlib.sha256(final_bytes).hexdigest() != final_sha:
        raise ValueError("original final protocol SHA mismatch")
    return notice_bytes, final_bytes, finals[0], latest


def _bid_summary(original: ET.Element) -> dict:
    apps = [e for e in original.iter() if _local(e) == "applicationInfo"]
    if not apps or len(apps) > 2000:
        raise ValueError("invalid original final number of bids")
    ids, admitted, rejected, ranked, amounts = set(), 0, 0, [], []
    missing_prices = 0
    for app in apps:
        number = _leaf(app, "appNumber")
        acceptance = _leaf(app, "admitted")
        rank = _leaf(app, "appRating")
        price = _leaf(app, "finalPrice")
        if not number or number in ids or acceptance not in {"true", "false"}:
            raise ValueError("invalid or ambiguous original final decision")
        ids.add(number)
        if acceptance == "false":
            rejected += 1
            if rank:
                raise ValueError("rejected bid improperly ranked")
            continue
        admitted += 1
        if not rank or not rank.isdigit():
            raise ValueError("admitted original rank missing")
        ranked.append(int(rank))
        if price is None:
            missing_prices += 1
        else:
            value = Decimal(price)
            if not value.is_finite() or value < 0:
                raise ValueError("invalid original price")
            amounts.append(value)
    if sorted(ranked) != list(range(1, admitted + 1)):
        raise ValueError("original admitted ranks not contiguous")
    if missing_prices and amounts:
        raise ValueError("partial finalPrices are not a comparable result")
    if amounts:
        return {
            "source_price_field_state": "VERIFIED_ADMITTED_FINAL_PRICE",
            "admitted_price_low_rub": str(min(amounts)),
            "admitted_price_median_rub": str(median(amounts)),
            "admitted_price_high_rub": str(max(amounts)),
            "extreme_sub_ruble_bid_detected": any(0 < p < Decimal("1.00") for p in amounts),
            "admitted_prices_for_individual_original_only": len(amounts),
            "admitted": admitted,
            "rejected": rejected,
            "total": len(apps),
        }
    return {
        "source_price_field_state": "MISSING_FINAL_PRICE_IN_ORIGINAL_EOK_PROTOCOL",
        "admitted_price_low_rub": None,
        "admitted_price_median_rub": None,
        "admitted_price_high_rub": None,
        "extreme_sub_ruble_bid_detected": False,
        "admitted_prices_for_individual_original_only": 0,
        "admitted": admitted,
        "rejected": rejected,
        "total": len(apps),
    }


def audit_case(pin: Pin, *, eis_root: Path, technical_scope_root: Path) -> dict:
    if not re.fullmatch(r"\d{19}", pin.number):
        raise ValueError("invalid pinned registry")
    archive_path = eis_root / pin.number / "documentation-archive.zip"
    scope_path = technical_scope_root / pin.number / "technical-scope.docx"
    if any(x.is_symlink() or not x.is_file() or not 100 < x.stat().st_size <= 8_000_000 for x in (archive_path, scope_path)):
        raise ValueError("missing/unsafe original file")
    archive_bytes = archive_path.read_bytes()
    scope_bytes = scope_path.read_bytes()
    if hashlib.sha256(archive_bytes).hexdigest() != pin.zip_sha256:
        raise ValueError("official EIS ZIP identity mismatch")
    if hashlib.sha256(scope_bytes).hexdigest() != pin.scope_sha256:
        raise ValueError("original filestore DOCX SHA mismatch")
    with ZipFile(io.BytesIO(archive_bytes)) as archive:
        notice, final, member, version = _notice_and_final(archive, pin.number, pin.final_sha256)
    root = _xml(notice)
    titles = [" ".join((x.text or "").split()) for x in root.iter() if _local(x) == "purchaseObjectInfo" and (x.text or "").strip()]
    if len(titles) != 1 or _normalized(pin.exact_scope_subject_fragment) not in _normalized(titles[0]):
        raise ValueError("original notice subject mismatch")
    numbers = {(x.text or "").strip() for x in root.iter() if _local(x) == "purchaseNumber" and re.fullmatch(r"\d{19}", (x.text or "").strip())}
    if numbers != {pin.number}:
        raise ValueError("notice registry mismatch")
    try:
        doc = Document(io.BytesIO(scope_bytes))
        lines = [p.text for p in doc.paragraphs]
        lines.extend(c.text for t in doc.tables for r in t.rows for c in r.cells)
        text = _normalized(" ".join(lines))
    except Exception as exc:
        raise ValueError("unable to extract original Word scope") from exc
    missing = [phrase for phrase in pin.scope_evidence_phrases if _normalized(phrase) not in text]
    if missing:
        raise ValueError("original scope mandatory anchor drift")
    bids = _bid_summary(_xml(final))
    return {
        "registry_number": pin.number,
        "original_notice_subject": titles[0],
        "official_zip_sha256": pin.zip_sha256,
        "original_notice_xml_sha256": hashlib.sha256(notice).hexdigest(),
        "original_final_xml_sha256": pin.final_sha256,
        "original_technical_scope_docx_sha256": pin.scope_sha256,
        "notice_revision": version,
        "original_final_member_type": member.split("_", 1)[0],
        "source_scope_anchors_verified": list(pin.scope_evidence_phrases),
        "scope_comparison_state": pin.comparison,
        "scope_comparison_reason": pin.reason,
        "accepted_to_aksay_normalized_basket": False,
        **bids,
    }


def audit_cohort(*, eis_root: Path, scope_root: Path) -> dict:
    cases = [audit_case(pin, eis_root=eis_root, technical_scope_root=scope_root) for pin in PINS]
    if sum(x["source_price_field_state"] == "VERIFIED_ADMITTED_FINAL_PRICE" for x in cases) != 5:
        raise ValueError("unexpected missing historical source price fields")
    if [x["total"] for x in cases] != [55, 6, 30, 28, 10, 11]:
        raise ValueError("historical source original application count changed")
    return {
        "schema_version": "ta-price01-official-website-original-scope-cohort-v1",
        "source_scope": "exact_source_44fz_archival_only",
        "original_verified_procedures": len(cases),
        "original_procedures_with_admitted_final_prices": 5,
        "original_procedures_without_final_price": 1,
        "related_but_nonpooled_additional_procedures": 2,
        "excluded_scope_mismatch_additional_procedures": 3,
        "new_admitted_same_basket_procurements": 0,
        "aksay_website_same_basket_count": 1,
        "profit_or_win_probability_forecast": False,
        "autonomous_bid_or_payment": False,
        "cases": cases,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--eis-root", type=Path, required=True)
    parser.add_argument("--scope-root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    report = audit_cohort(eis_root=args.eis_root, scope_root=args.scope_root)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
    print("SOURCE_CASES",report["original_verified_procedures"],
          "VERIFIED_FINAL_PRICES",report["original_procedures_with_admitted_final_prices"],
          "AKSAY_COMPARABLE",report["aksay_website_same_basket_count"])


if __name__ == "__main__":
    main()
