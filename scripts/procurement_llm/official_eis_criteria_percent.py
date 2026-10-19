"""Verified, source-SHA-bound percentage values from a real criteria DOCX."""
from __future__ import annotations

import hashlib
import io
import json
import os
import stat
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from sqlalchemy.orm import Session

from scripts.procurement_llm.analysis_queue_worker import PinnedSource, attest_source
from scripts.procurement_llm.collect_public_market_originals import OWNER_PRIVATE
from scripts.procurement_llm.prepare_eis_original_documents import (
    SourcePreparationError,
)
from scripts.procurement_llm.source_derived_docx import canonical_docx_bytes
from src.tender_research.repository import TenderRepository

NUMBER = "0163200000326005243"
PINNED_CRITERIA_SHA256 = "a712e03f43f5967a9bf42e1c0589b670769af352aea3a32e18c96d416724583d"
NAME = "Критерии оценки заявок и порядок рассмотрения.docx"
SUFFIX = "-derived-criteria-percent"
SPECS = [
    ("Цена контракта", "30", 1, 2, "цен"),
    ("Квалификация участников", "70", 2, 2, "квалификац"),
    ("Цена исполненных договоров", "80", 2, 6, "общая цена"),
    ("Количество исполненных договоров", "20", 3, 6, "общее количество"),
]
HEADERS = {
    2: "Значимость критерия оценки, процентов",
    6: "Значимость показателя, детализирующего показатель оценки, процентов",
}


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_verified_facts(original: bytes) -> list[dict]:
    if not 0 < len(original) < 3_000_000:
        raise SourcePreparationError("criteria source size invalid")
    try:
        doc = Document(io.BytesIO(original))
    except (ValueError, KeyError, OSError, BadZipFile) as exc:
        raise SourcePreparationError("criteria original invalid DOCX") from exc
    matches = [(i, t) for i, t in enumerate(doc.tables)
               if len(t.rows) >= 4 and len(t.columns) == 8
               and all(" ".join(t.cell(0, k).text.split()) == v
                       for k, v in HEADERS.items())]
    if len(matches) != 1:
        raise SourcePreparationError("criteria percent source table ambiguous")
    table_idx, table = matches[0]
    facts = []
    for meaning, expected, row, col, keyword in SPECS:
        value = " ".join(table.cell(row, col).text.split())
        label_col = 1 if col == 2 else 5
        label = " ".join(table.cell(row, label_col).text.split())
        if value != expected or keyword not in label.casefold():
            raise SourcePreparationError("source percentage/header/label mismatch")
        facts.append({
            "meaning": meaning, "value_percent": int(value),
            "table": table_idx + 1, "row": row + 1, "column": col + 1,
            "source_header": HEADERS[col], "source_label": label,
        })
    if sum(f["value_percent"] for f in facts[:2]) != 100 or sum(
        f["value_percent"] for f in facts[2:4]
    ) != 100:
        raise SourcePreparationError("source criteria weight sum invalid")
    return facts


def context_blob(src: PinnedSource, *, root: Path = OWNER_PRIVATE) -> tuple[bytes, dict]:
    if src.procurement_ref != NUMBER or src.platform != "eis":
        raise SourcePreparationError("Smolensk source-pinned criteria only")
    attest_source(src, root=root)
    folder = root / "eis-source-documents" / src.procurement_ref / src.original_sha256
    manifest_path = folder / "source-manifest.json"
    if (manifest_path.is_symlink() or not manifest_path.is_file()
            or stat.S_IMODE(manifest_path.stat().st_mode) != 0o600):
        raise SourcePreparationError("source manifest missing/unsafe")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if (manifest.get("archived_original_zip_sha256") != src.original_sha256
            or manifest.get("source_observation_id") != src.observation_id
            or manifest.get("all_original_binaries_saved") is not True):
        raise SourcePreparationError("source manifest mismatch")
    entries = [e for e in manifest["attachments"]
               if e["original_name"] == NAME and e["binary_type"] == "DOCX_MAIN_BODY_OK"]
    if len(entries) != 1:
        raise SourcePreparationError("criteria original ambiguous or missing")
    entry = entries[0]
    original = Path(entry["owner_private_original_path"])
    if (original.is_symlink() or not original.is_file()
            or original.parent.resolve() != folder.resolve()
            or stat.S_IMODE(original.stat().st_mode) != 0o600):
        raise SourcePreparationError("criteria source path unsafe")
    if entry["original_binary_sha256"] != PINNED_CRITERIA_SHA256:
        raise SourcePreparationError("pinned official criteria original SHA changed")
    binary = original.read_bytes()
    if sha(binary) != entry["original_binary_sha256"]:
        raise SourcePreparationError("criteria source content SHA conflict")
    facts = read_verified_facts(binary)
    generated = Document()
    generated.add_heading("ПРОИЗВОДНЫЕ СВЕДЕНИЯ О ПРОЦЕНТАХ ОЦЕНКИ, НЕ ОРИГИНАЛ ЕИС", 0)
    generated.add_paragraph("Оригинал DOCX SHA256: " + entry["original_binary_sha256"])
    generated.add_paragraph("Оригинал ZIP ЕИС SHA256: " + src.original_sha256)
    for fact in facts:
        generated.add_paragraph(
            f"{fact['meaning']}: {fact['value_percent']}%. "
            f"Из оригинала {NAME}: таблица {fact['table']}, строка {fact['row']}, "
            f"столбец {fact['column']}; заголовок «{fact['source_header']}»."
        )
    buffer = io.BytesIO()
    generated.save(buffer)
    result = canonical_docx_bytes(buffer.getvalue())
    return result, {
        "schema": "arvectum_eis_criteria_percent_from_original_cells_v1",
        "source_original_archive_sha256": src.original_sha256,
        "source_observation_id": src.observation_id,
        "criteria_original_sha256": entry["original_binary_sha256"],
        "facts": facts, "derived_docx_sha256": sha(result),
        "derived_not_official_signed": True,
        "live_status_verified": False,
        "human_training_gold": False,
    }


def store_criteria_context(session: Session, src: PinnedSource,
                           *, root: Path = OWNER_PRIVATE) -> dict:
    from scripts.procurement_llm.prepare_eis_original_documents import (
        verify_registered_exact_source_set,
    )
    result = verify_registered_exact_source_set(session, src.procurement_ref, root=root)
    if result["original_archive_sha256"] != src.original_sha256:
        raise SourcePreparationError("registry source changed")
    binary, receipt = context_blob(src, root=root)
    folder = root / "eis-source-documents" / src.procurement_ref / src.original_sha256
    path = folder / ("derived-criteria-" + sha(binary) + ".docx")
    if path.is_symlink():
        raise SourcePreparationError("criteria output symlink")
    if path.exists():
        if (not path.is_file() or stat.S_IMODE(path.stat().st_mode) != 0o600
                or path.read_bytes() != binary):
            raise SourcePreparationError("criteria output not immutable")
    else:
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        with os.fdopen(fd, "wb") as stream:
            stream.write(binary);stream.flush();os.fsync(stream.fileno())
    tender = TenderRepository(session).get_tender_by_registry_number(src.procurement_ref)
    if tender is None:
        raise SourcePreparationError("registered original tender missing")
    TenderRepository(session).upsert_document({
        "tender_id": tender.id,
        "source_document_id": f"eis-archived-v2{SUFFIX}",
        "file_name": "Подтверждённые веса критериев ЕИС (производный).docx",
        "file_url": None, "local_path": str(path),
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": len(binary), "sha256": sha(binary),
        "download_status": "downloaded", "text_extraction_status": "pending",
        "raw_meta": receipt,
    })
    session.flush()
    return {"derived_sha256": sha(binary), "original_criteria_sha256":
            receipt["criteria_original_sha256"],
            "verified_percentages": [f["value_percent"] for f in receipt["facts"]]}
