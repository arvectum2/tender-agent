"""Reproducible 44-FZ official XML evidence for local Gemma RAG.

The signed-off original remains the precise owner-private getDocsIP ZIP.
This *derived*, transparent DOCX representation carries typed facts and the
complete XML object-characteristic text into Data Platform. No external LLM
or guessed data. Not a signed/current EIS document.
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import stat
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from sqlalchemy.orm import Session

from scripts.procurement_llm.analysis_queue_worker import PinnedSource, attest_source
from scripts.procurement_llm.collect_public_market_originals import OWNER_PRIVATE
from scripts.procurement_llm.prepare_eis_original_documents import (
    SourcePreparationError,
    _selected_notice,
    verify_registered_exact_source_set,
)
from scripts.procurement_llm.source_derived_docx import canonical_docx_bytes
from src.tender_research.repository import TenderRepository

MAX_CONTEXT_CHARS = 80000
LABEL = "ПРОИЗВОДНЫЙ ОБЗОР ИСХОДНОГО XML ЕИС — НЕ ОФИЦИАЛЬНЫЙ ПОДПИСАННЫЙ ФАЙЛ"


def _local(node: ET.Element) -> str:
    return node.tag.rsplit("}", 1)[-1]


def _direct(node: ET.Element, path: tuple[str, ...]) -> list[ET.Element]:
    nodes = [node]
    for segment in path:
        nodes = [child for parent in nodes for child in parent
                 if _local(child) == segment]
    return nodes


def _single(root: ET.Element, path: tuple[str, ...]) -> str:
    found = [" ".join((x.text or "").split()) for x in _direct(root, path)]
    if len(found) != 1 or not found[0]:
        raise SourcePreparationError("required official XML path missing/ambiguous")
    return found[0]


def render_source_facts(xml: bytes, *, registry_number: str,
                        xml_sha: str, archive_sha: str, revision: int) -> list[str]:
    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
        raise SourcePreparationError("XML entity expansion forbidden")
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        raise SourcePreparationError("invalid official XML") from exc
    candidates = [node for node in root.iter()
                  if _local(node) in {"epNotificationEOK2020", "epNotificationEF2020",
                                      "epNotificationEZK2020"}]
    if len(candidates) != 1:
        raise SourcePreparationError("ambiguous official EIS notice type")
    notice = candidates[0]
    fields = {
        "Номер извещения": ("commonInfo", "purchaseNumber"),
        "Название услуги": ("commonInfo", "purchaseObjectInfo"),
        "Дата публикации извещения (архивная)": ("commonInfo", "publishDTInEIS"),
        "Способ закупки": ("commonInfo", "placingWay", "name"),
        "Электронная площадка": ("commonInfo", "ETP", "name"),
        "НМЦК РУБ — сумма контракта, НЕ лимит квалификационного опыта": (
            "notificationInfo", "contractConditionsInfo", "maxPriceInfo", "maxPrice"),
        "Начало приёма заявок": (
            "notificationInfo", "procedureInfo", "collectingInfo", "startDT"),
        "Окончание приёма заявок": (
            "notificationInfo", "procedureInfo", "collectingInfo", "endDT"),
    }
    out = [
        LABEL,
        "Источник: официальный getDocsIP XML внутри проверенного ZIP ЕИС.",
        f"SHA256 ZIP: {archive_sha}",
        f"SHA256 XML: {xml_sha}",
        f"Архивная версия: {revision}",
        "Важно: это снимок исходной редакции, НЕ подтверждает её актуальность сегодня.",
    ]
    for label, path in fields.items():
        value = _single(notice, path)
        if len(value) > 1000:
            raise SourcePreparationError("unexpected oversized official source scalar")
        if label == "Номер извещения" and value != registry_number:
            raise SourcePreparationError("official XML purchase number differs")
        if label.startswith("НМЦК") and (
            not re.fullmatch(r"[0-9]{1,13}\.[0-9]{1,2}", value)
            or float(value) <= 0
        ):
            raise SourcePreparationError("official EIS NMC invalid")
        out.append(f"{label}: {value}")
    characteristics = [
        x for x in notice.iter() if _local(x) == "characteristicsUsingTextForm"
    ]
    if not 1 <= len(characteristics) <= 100:
        raise SourcePreparationError("official object characteristics missing/oversized")
    out.append(f"Характеристик объекта закупки в исходном XML: {len(characteristics)}")
    for index, element in enumerate(characteristics, 1):
        value = _single(element, ("name",))
        # The official EIS value may carry the entire paragraph text; no truncation.
        if len(value) > 10000:
            raise SourcePreparationError("one official object characteristic oversized")
        out.append(f"Требование ТЗ #{index:02d}: {value}")
    if sum(len(x) for x in out) > MAX_CONTEXT_CHARS:
        raise SourcePreparationError("derived XML contextual document oversized")
    return out


def source_context_blob(src: PinnedSource, *, root: Path = OWNER_PRIVATE) -> tuple[bytes, dict]:
    attest_source(src, root=root)
    xml, info = _selected_notice(src)
    fields = render_source_facts(
        xml, registry_number=src.procurement_ref,
        xml_sha=info["notice_xml_sha256"],
        archive_sha=src.original_sha256, revision=info["source_version"],
    )
    doc = Document()
    doc.add_heading(LABEL, 0)
    for index, row in enumerate(fields):
        if row.startswith("Требование ТЗ #"):
            doc.add_heading(row.split(":", 1)[0], level=2)
            doc.add_paragraph(row.split(":", 1)[1].lstrip())
        elif index != 0:
            doc.add_paragraph(row)
    buffer = io.BytesIO()
    doc.save(buffer)
    raw = canonical_docx_bytes(buffer.getvalue())
    if not 0 < len(raw) <= 400000:
        raise SourcePreparationError("derived XML DOCX out of bound")
    return raw, {
        "schema": "official_eis_archived_xml_derived_full_text_v1",
        "purchase_number": src.procurement_ref,
        "source_observation_id": src.observation_id,
        "source_original_zip_sha256": src.original_sha256,
        "source_notice_xml_sha256": info["notice_xml_sha256"],
        "archived_notice_revision": info["source_version"],
        "characteristics_included": sum(x.startswith("Требование ТЗ #") for x in fields),
        "source_facts": {x.split(": ", 1)[0]: x.split(": ", 1)[1]
                         for x in fields if ": " in x and not x.startswith("Требование ТЗ #")},
        "derived_docx_sha256": hashlib.sha256(raw).hexdigest(),
        "not_original_eis_signed_document": True,
        "current_legal_status_verified": False,
    }


def store_xml_derived_context(session: Session, src: PinnedSource,
                              *, root: Path = OWNER_PRIVATE) -> dict:
    proof = verify_registered_exact_source_set(session, src.procurement_ref, root=root)
    if proof["original_archive_sha256"] != src.original_sha256:
        raise SourcePreparationError("base document set not exact ZIP SHA")
    data, receipt = source_context_blob(src, root=root)
    directory = root / "eis-source-documents" / src.procurement_ref / src.original_sha256
    if directory.is_symlink() or not directory.is_dir():
        raise SourcePreparationError("expected owned original folder missing")
    dest = directory / ("derived-xml-v" + str(receipt["archived_notice_revision"])
                        + "-" + receipt["derived_docx_sha256"] + ".docx")
    if dest.is_symlink():
        raise SourcePreparationError("derived XML evidence symlink forbidden")
    if dest.exists():
        if (not dest.is_file() or stat.S_IMODE(dest.stat().st_mode) != 0o600
                or hashlib.sha256(dest.read_bytes()).hexdigest()
                != receipt["derived_docx_sha256"]):
            raise SourcePreparationError("modified or nonprivate derived XML document")
    else:
        fd = os.open(dest, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        with os.fdopen(fd, "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
    repo = TenderRepository(session)
    tender = repo.get_tender_by_registry_number(src.procurement_ref)
    if tender is None:
        raise SourcePreparationError("source-bound Tender Research row missing")
    repo.upsert_document({
        "tender_id": tender.id,
        "source_document_id": (
            "eis-archived-v" + str(receipt["archived_notice_revision"])
            + "-derived-full-xml"),
        "file_name": "Проверенные сведения XML ЕИС v"
                     + str(receipt["archived_notice_revision"]) + " (производный).docx",
        "file_url": None, "local_path": str(dest),
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": len(data),
        "sha256": receipt["derived_docx_sha256"],
        "download_status": "downloaded",
        "text_extraction_status": "pending",
        "raw_meta": receipt,
    })
    session.flush()
    return {
        "source_archive_sha256": src.original_sha256,
        "source_xml_sha256": receipt["source_notice_xml_sha256"],
        "docx_sha256": receipt["derived_docx_sha256"],
        "full_xml_characteristics": receipt["characteristics_included"],
        "notice_nmck_rub": receipt["source_facts"]["НМЦК РУБ — сумма контракта, НЕ лимит квалификационного опыта"],
        "notice_submission_end": receipt["source_facts"]["Окончание приёма заявок"],
        "derived_source_document": str(dest),
        "no_new_original_facts_invented": True,
        "current_legal_status_verified": False,
    }
