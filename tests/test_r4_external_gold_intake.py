"""Synthetic external-reviewer signatures, not actual procurement gold.

All keys here are throwaway test-only; NO real external expert is enrolled.
"""
from __future__ import annotations

import base64
import copy
import hashlib
from pathlib import Path

import pytest
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from docx import Document

from src.modules.benchmark_pipeline.contract import (
    BenchmarkContractError,
    source_bundle_sha256,
)
from src.modules.benchmark_pipeline.workflow import prepare_evaluator_bundle
from src.modules.tender_operator_agent_demo.r4_external_gold_intake import (
    attestation_payload,
    freeze_verified_external_gold,
    verify_independent_expert_submission,
)

P0 = "2026-10-14T17:00:00+00:00"
P1 = "2026-10-14T18:00:00+00:00"
P2 = "2026-10-14T19:00:00+00:00"
P3 = "2026-10-14T20:00:00+00:00"
CLAUSE = "Подрядчик должен передать исходный код заказчику по окончании работ."


def synthetic_signed_review(tmp_path: Path):
    (tmp_path / "sources").mkdir()
    path = tmp_path / "sources" / "example.docx"
    doc = Document()
    doc.add_paragraph("Тестовый договор о разработке программного обеспечения.")
    doc.add_paragraph(CLAUSE)
    doc.save(path)
    docs = [{
        "path": "sources/example.docx",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_url": "https://example.org/official/eis/123",
    }]
    manifest = {
        "schema_version": "1.1.0",
        "case_id": "SYNTHETIC-INDEPENDENT-REVIEW",
        "procurement": {"notice_number": "9999999999999999999"},
        "source_urls": ["https://example.org/official/eis/123"],
        "acquired_at": P0,
        "documents": docs,
        "source_scope": "Synthetic source, not a real EIS procurement",
        "source_bundle_sha256": source_bundle_sha256(docs),
        "source_conflict": False,
        "provenance_sufficient": True,
    }
    bundle = prepare_evaluator_bundle(manifest, prepared_at=P0)
    discovery = {
        "schema_version": "1.1.0",
        "case_id": bundle["case_id"],
        "source_bundle_sha256": bundle["source_bundle_sha256"],
        "label": "RELEVANT",
        "reason": "Synthetic offline evaluation.",
        "confidence": .8,
        "evidence": [{
            "source_ref": docs[0]["path"],
            "locator": "PARAGRAPH-0002",
            "quote": CLAUSE,
        }],
        "evaluator": "outside-test-reviewer",
        "evaluated_at": P1,
    }
    truth = {
        "schema_version": "1.1.0",
        "case_id": bundle["case_id"],
        "source_bundle_sha256": bundle["source_bundle_sha256"],
        "facts": [{
            "field": "ip.source_code_handover",
            "value": "Передать исходный код заказчику",
            "evidence": [{
                "source_ref": docs[0]["path"],
                "locator": "PARAGRAPH-0002",
                "quote": CLAUSE,
            }],
            "confidence": .8,
            "abstention": "ASSERTED",
            "materiality": "MATERIAL",
        }],
        "confidence": .8,
        "evaluator": "outside-test-reviewer",
        "evaluated_at": P1,
    }
    key = Ed25519PrivateKey.generate()
    raw_public = key.public_key().public_bytes(
        encoding=serialization.Encoding.Raw,
        format=serialization.PublicFormat.Raw,
    )
    keys = {
        "outside-test-reviewer": {
            "role": "QUALIFIED_INDEPENDENT_LEGAL_EXPERT",
            "owner_identity_verified": True,
            "ed25519_public_key_b64": base64.b64encode(raw_public).decode(),
        }
    }
    signature = key.sign(attestation_payload(
        bundle, discovery, truth,
        reviewer_id="outside-test-reviewer", attested_at=P2,
    ))
    attestation = {
        "reviewer_id": "outside-test-reviewer",
        "attested_at": P2,
        "no_prior_sut_access": True,
        "ed25519_signature_b64": base64.b64encode(signature).decode(),
    }
    args = {
        "evaluator_bundle": bundle,
        "discovery_label": discovery,
        "document_truth": truth,
        "source_manifest": manifest,
        "original_source_root": tmp_path,
        "signed_attestation": attestation,
        "trusted_reviewer_keys": keys,
    }
    return args


def test_signed_source_grounded_external_review_can_freeze_synthetic_only(tmp_path):
    a = synthetic_signed_review(tmp_path)
    receipt = verify_independent_expert_submission(**a)
    assert receipt["status"] == "SIGNATURE_AND_SOURCE_QUOTE_PRECHECK_PASSED"
    assert not receipt["independent_legal_opinion_issued_by_this_tool"]
    frozen = freeze_verified_external_gold(**a, frozen_at=P3)
    assert frozen["source_bundle_sha256"] == a["evaluator_bundle"]["source_bundle_sha256"]
    assert len(frozen["label_set_sha256"]) == 64


@pytest.mark.parametrize("modification", [
    lambda a: a["trusted_reviewer_keys"].clear(),
    lambda a: a["trusted_reviewer_keys"]["outside-test-reviewer"].update(owner_identity_verified=False),
    lambda a: a["signed_attestation"].update(no_prior_sut_access=False),
    lambda a: a["signed_attestation"].update(attested_at=P0),
    lambda a: a["signed_attestation"].update(ed25519_signature_b64=base64.b64encode(b"Z"*64).decode()),
    lambda a: a["document_truth"]["facts"][0].update(value="Подделано"),
    lambda a: a["document_truth"]["facts"][0]["evidence"][0].update(quote="Фантазийный текст договорного обязательства"),
    lambda a: a["discovery_label"]["evidence"][0].update(quote="Выгаданная причина релевантности данной закупки"),
    lambda a: a["document_truth"]["facts"][0]["evidence"][0].update(source_ref="https://example.org/official/eis/123"),
    lambda a: a["document_truth"]["facts"][0]["evidence"][0].update(locator=None),
    lambda a: a["document_truth"]["facts"][0]["evidence"][0].update(locator="PARAGRAPH-9999"),
])
def test_missing_trust_or_forged_source_or_altered_label_rejected(tmp_path, modification):
    args = synthetic_signed_review(tmp_path)
    modification(args)
    with pytest.raises((BenchmarkContractError, ValueError)):
        freeze_verified_external_gold(**args, frozen_at=P3)


def test_modified_original_or_early_freeze_rejected(tmp_path):
    args = synthetic_signed_review(tmp_path)
    with pytest.raises(BenchmarkContractError, match="strictly after"):
        freeze_verified_external_gold(**args, frozen_at=P2)
    doc_path = tmp_path / "sources" / "example.docx"
    doc_path.write_bytes(b"replaced raw bytes")
    with pytest.raises(BenchmarkContractError, match="hash mismatch"):
        freeze_verified_external_gold(**args, frozen_at=P3)


def test_no_actual_holdout_gold_is_created_by_validation(tmp_path):
    args = synthetic_signed_review(tmp_path)
    shadow = copy.deepcopy(args)
    assert not list(tmp_path.glob("**/blind_document_truth.json"))
    assert not list(tmp_path.glob("**/frozen_label.json"))
    assert shadow["document_truth"]["case_id"].startswith("SYNTHETIC")
