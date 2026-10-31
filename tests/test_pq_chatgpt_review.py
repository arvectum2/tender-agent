"""ChatGPT classifications must be tied to exact original text."""
import pytest

from src.modules.market_intelligence.pq_chatgpt_review import validate_review


def test_grounded_review():
    packet = {"registry_number": "32313113128", "documents": [
        {"sha256": "a" * 64, "text": "поставка автомобиля для аварийно-спасательного формирования"}]}
    review = {"registry_number": "32313113128", "classification": "irrelevant",
              "evidence": [{"sha256": "a" * 64,
                            "quote": "поставка автомобиля для аварийно-спасательного формирования"}]}
    result = validate_review(packet=packet, review=review)
    assert result["classification"] == "irrelevant"
    assert result["market_total_rub"] is None
    review["evidence"][0]["quote"] = "несуществующая поставка компенсаторов мощности"
    with pytest.raises(ValueError):
        validate_review(packet=packet, review=review)


def test_legacy_doc_original_sha_and_quoted_text(tmp_path, monkeypatch):
    from unittest.mock import Mock

    from src.modules.market_intelligence import pq_chatgpt_review as mod
    source = tmp_path / "technical.doc"
    source.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"original-body")
    monkeypatch.setattr(mod.sys, "platform", "darwin")
    seen = []

    def run(command, **kwargs):
        seen.append((command, kwargs))
        return Mock(returncode=0, stdout="УКРМ 10 кВ, 150 квар".encode())

    monkeypatch.setattr(mod.subprocess, "run", run)
    text, method = mod._extract(source)
    assert text == "УКРМ 10 кВ, 150 квар"
    assert method == "legacy_doc_text"
    assert seen[0][0] == ["/usr/bin/textutil", "-convert", "txt", "-stdout", str(source)]
    assert seen[0][1]["timeout"] == 25


def test_legacy_doc_rejects_non_cfb_without_running(tmp_path, monkeypatch):
    from src.modules.market_intelligence import pq_chatgpt_review as mod
    source = tmp_path / "fake.doc"
    source.write_bytes(b"not a binary Word file")
    monkeypatch.setattr(mod.subprocess, "run", lambda *_a, **_kw: pytest.fail("executed"))
    assert mod._extract(source) == ("", "invalid_legacy_doc_container")


def test_legacy_doc_missing_linux_extractor_is_provisional(tmp_path, monkeypatch):
    from src.modules.market_intelligence import pq_chatgpt_review as mod
    source = tmp_path / "technical.doc"
    source.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"body")
    monkeypatch.setattr(mod.sys, "platform", "linux")
    monkeypatch.setattr(mod.shutil, "which", lambda _x: None)
    assert mod._extract(source) == ("", "legacy_doc_extractor_unavailable")


def test_legacy_doc_packet_binds_original_bytes(tmp_path, monkeypatch):
    import hashlib
    from unittest.mock import Mock

    from src.modules.market_intelligence import pq_chatgpt_review as mod

    original = bytes.fromhex("d0cf11e0a1b11ae1") + b"original-body"
    (tmp_path / "spec.doc").write_bytes(original)
    ontology = tmp_path / "ontology.yaml"
    ontology.write_text("families: []", encoding="utf-8")
    monkeypatch.setattr(mod.sys, "platform", "darwin")
    monkeypatch.setattr(
        mod.subprocess, "run",
        lambda *_a, **_kw: Mock(returncode=0, stdout="УКРМ 10 кВ".encode()),
    )
    packet = mod.prepare_packet(
        document_dir=tmp_path, ontology_path=ontology, registry_number="32413764701",
    )
    assert len(packet["documents"]) == 1
    assert packet["documents"][0]["sha256"] == hashlib.sha256(original).hexdigest()
    assert packet["documents"][0]["text"] == "УКРМ 10 кВ"


def test_xlsx_late_pq_cell_retains_original_sha_and_quote(tmp_path):
    import hashlib

    from openpyxl import Workbook

    from src.modules.market_intelligence import pq_chatgpt_review as mod

    book = Workbook()
    sheet = book.active
    for i in range(1, 2001):
        sheet.cell(i, 1, "Радиотехническое оборудование " + "тест " * 15)
    quote = "КТП 10/0,4 кВ; в составе РУНН УКРМ-2шт; напряжение УКРМ не указано"
    sheet.cell(2001, 2, quote)
    source = tmp_path / "original.xlsx"
    book.save(source)
    ontology = tmp_path / "ontology.yaml"
    ontology.write_text("families: []", encoding="utf-8")
    packet = mod.prepare_packet(
        document_dir=tmp_path, ontology_path=ontology, registry_number="0548100001225000001",
    )
    doc = packet["documents"][0]
    assert doc["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert quote in doc["text"]
    assert "ROW 2001 COL 2" in doc["text"]
    assert len(doc["text"]) <= mod.MAX_TEXT
    review = {"registry_number": "0548100001225000001", "classification": "uncertain",
              "evidence": [{"sha256": doc["sha256"], "quote": quote}]}
    assert mod.validate_review(packet=packet, review=review)["classification"] == "uncertain"
