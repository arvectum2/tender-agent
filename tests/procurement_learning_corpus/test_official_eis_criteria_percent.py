"""Original source cell-based EIS significance facts must not be guessed."""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from scripts.procurement_llm import official_eis_criteria_percent as p
from scripts.procurement_llm.prepare_eis_original_documents import (
    SourcePreparationError,
)


def table_bytes(*, corrupt: str | None = None, extra: bool = False) -> bytes:
    doc = Document()
    table = doc.add_table(rows=4, cols=8)
    for col, header in p.HEADERS.items():
        table.cell(0, col).text = header
    for i, (_, value, row, col, label) in enumerate(p.SPECS):
        table.cell(row, col).text = value if corrupt != f"value{i}" else "99"
        label_col = 1 if col == 2 else 5
        table.cell(row, label_col).text = label
    if corrupt == "header":
        table.cell(0, 2).text = "Значимость критерия оценки"
    if extra:
        copy = doc.add_table(rows=4, cols=8)
        for col, header in p.HEADERS.items():
            copy.cell(0, col).text = header
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_criteria_realistic_table_values_and_source_coordinates():
    facts = p.read_verified_facts(table_bytes())
    assert [x["value_percent"] for x in facts] == [30, 70, 80, 20]
    assert [x["column"] for x in facts] == [3, 3, 7, 7]
    assert all("процентов" in x["source_header"] for x in facts)


@pytest.mark.parametrize("modified", ["value0", "value1", "value2", "value3", "header"])
def test_criteria_no_inferred_percentage_when_source_modified(modified):
    with pytest.raises(SourcePreparationError):
        p.read_verified_facts(table_bytes(corrupt=modified))


def test_duplicate_criteria_table_does_not_silently_choose_first():
    with pytest.raises(SourcePreparationError, match="ambiguous"):
        p.read_verified_facts(table_bytes(extra=True))


def test_bad_non_docx_file_is_not_accepted_as_criteria():
    with pytest.raises(SourcePreparationError):
        p.read_verified_facts(b"not a valid original DOCX")



def test_scoped_to_exact_smolenks_original_and_not_unrelated_tenders(tmp_path):
    from scripts.procurement_llm.analysis_queue_worker import PinnedSource

    unknown=PinnedSource(
        observation_id="a"*64,procurement_ref="0332100025426000039",
        event_kind="manual:ta-eis-zip:sha",original_sha256="b"*64,
        path=tmp_path/"unrelated.zip",platform="eis",
    )
    with pytest.raises(SourcePreparationError,match="Smolensk"):
        p.context_blob(unknown,root=tmp_path)
