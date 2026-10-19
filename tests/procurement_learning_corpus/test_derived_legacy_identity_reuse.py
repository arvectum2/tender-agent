"""Legacy derived DOCX re-import is SHA-preserving, not an implicit revision."""
from __future__ import annotations

import hashlib
import io
from types import SimpleNamespace
from zipfile import ZipFile, ZipInfo

import pytest
from docx import Document

from scripts.procurement_llm.source_derived_docx import (
    canonical_docx_bytes,
    reuse_existing_verified_source_document,
)


def _original_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Official original XML: 497178.33")
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def _different_container_same_content(raw: bytes) -> bytes:
    result = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as original, ZipFile(result, "w") as altered:
        for entry in original.infolist():
            info = ZipInfo(entry.filename, date_time=(2026, 10, 19, 10, 0, 0))
            info.compress_type = entry.compress_type
            altered.writestr(info, original.read(entry.filename))
    return result.getvalue()


def test_old_physical_sha_reused_if_identical_ooxml_members(tmp_path):
    new_bytes = canonical_docx_bytes(_original_docx())
    legacy = _different_container_same_content(new_bytes)
    assert legacy != new_bytes
    path = tmp_path / "old-derived.docx"
    path.write_bytes(legacy)
    path.chmod(0o600)
    source_sha = hashlib.sha256(legacy).hexdigest()
    receipt = {
        "source_original_zip_sha256": "a" * 64,
        "source_notice_xml_sha256": "b" * 64,
        "derived_docx_sha256": hashlib.sha256(new_bytes).hexdigest(),
    }
    doc = SimpleNamespace(
        source_document_id="eis-archived-v2-derived-full-xml",
        local_path=str(path),
        sha256=source_sha,
        raw_meta={**receipt, "derived_docx_sha256": source_sha},
    )
    tender = SimpleNamespace(documents=[doc])
    result = reuse_existing_verified_source_document(
        tender, doc.source_document_id, new_bytes, receipt, tmp_path,
    )
    assert result == (str(path), source_sha)
    assert path.read_bytes() == legacy


def test_old_source_rejected_for_different_content(tmp_path):
    old = _original_docx()
    path = tmp_path / "old.docx"
    path.write_bytes(old)
    path.chmod(0o600)
    old_sha = hashlib.sha256(old).hexdigest()
    receipt = {"source_original_zip_sha256": "a" * 64, "derived_docx_sha256": old_sha}
    document = SimpleNamespace(
        source_document_id="stable", local_path=str(path), sha256=old_sha,
        raw_meta=receipt,
    )
    different = Document()
    different.add_paragraph("Official original XML: 999999.99")
    out = io.BytesIO()
    different.save(out)
    with pytest.raises(ValueError, match="content conflict"):
        reuse_existing_verified_source_document(
            SimpleNamespace(documents=[document]), "stable",
            canonical_docx_bytes(out.getvalue()), receipt, tmp_path,
        )


def test_legacy_corruption_fails_even_if_expected_text_matches(tmp_path):
    raw = _original_docx()
    path = tmp_path / "old.docx"
    path.write_bytes(raw)
    path.chmod(0o600)
    doc = SimpleNamespace(
        source_document_id="stable", local_path=str(path), sha256="b" * 64,
        raw_meta={"derived_docx_sha256": "b" * 64},
    )
    with pytest.raises(ValueError):
        reuse_existing_verified_source_document(
            SimpleNamespace(documents=[doc]), "stable",
            canonical_docx_bytes(raw), {"derived_docx_sha256": "a" * 64}, tmp_path,
        )
