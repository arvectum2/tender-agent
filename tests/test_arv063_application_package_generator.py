from __future__ import annotations

import hashlib
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.modules.document_store.models import DocumentArtifact


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _deal(client, suffix: str = "") -> str:
    response = client.post(
        "/deals",
        json={
            "title": f"ARV-063 {suffix}",
            "initial_source_type": "manual_entry",
            "direction_type": "SUPPLY",
            "domain_type": "goods",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["deal_id"]


def _artifact(client, deal_id: str, root: Path, relative: str) -> str:
    path = root / relative
    response = client.post(
        "/artifacts",
        json={
            "deal_id": deal_id,
            "artifact_type": "OTHER",
            "file_name": path.name,
            "mime_type": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                if path.suffix.lower() == ".xlsx"
                else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                if path.suffix.lower() == ".docx"
                else "application/octet-stream"
            ),
            "storage_uri": relative,
            "checksum_sha256": _sha(path),
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["artifact_ref"]


def test_arv063_xlsx_working_kp_preserves_layout_and_versions(
    client, session, tmp_path: Path, monkeypatch
):
    monkeypatch.setattr(
        "src.modules.application_package_generator.service.load_config",
        lambda: SimpleNamespace(data_dir=str(tmp_path)),
    )
    deal_id = _deal(client, "xlsx")
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    template = template_dir / "kp-template.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "КП"
    sheet["A1"] = "Коммерческое предложение"
    sheet["A1"].font = Font(bold=True)
    sheet["B2"] = "PRICE_PLACEHOLDER"
    sheet["C5"] = "=SUM(1,2)"
    sheet.column_dimensions["A"].width = 31
    sheet.merge_cells("D1:E1")
    sheet["D1"] = "Не менять"
    workbook.save(template)

    template_ref = _artifact(client, deal_id, tmp_path, "templates/kp-template.xlsx")
    payload = {
        "deal_id": deal_id,
        "template_artifact_ref": template_ref,
        "expected_template_version": 1,
        "document_role": "COMMERCIAL_PROPOSAL",
        "fields": [
            {
                "target_locator": "xlsx:cell:КП!B2",
                "value": 125000,
                "source_type": "DEAL",
                "source_ref": deal_id,
            }
        ],
    }
    first = client.post("/api/application-drafts/generate", json=payload)
    assert first.status_code == 201, first.text
    one = first.json()
    assert one["output_file_name"] == "working_kp.xlsx"
    assert one["generation_version_no"] == 1
    assert one["review_state"] == "DRAFT_REVIEW_ONLY"
    assert one["signature_performed"] is False
    assert one["submission_performed"] is False
    assert one["external_delivery_performed"] is False
    assert one["template_sha256"] == _sha(template)
    assert one["field_provenance"] == [
        {
            "source_type": "DEAL",
            "source_ref": deal_id,
            "target_locator": "xlsx:cell:КП!B2",
            "original_value": "PRICE_PLACEHOLDER",
            "generated_value": 125000,
        }
    ]

    first_path = tmp_path / one["output_storage_uri"]
    assert first_path.is_file()
    rendered = load_workbook(first_path, data_only=False)
    rs = rendered["КП"]
    assert rs["B2"].value == 125000
    assert rs["A1"].value == "Коммерческое предложение"
    assert rs["A1"].font.bold is True
    assert rs["C5"].value == "=SUM(1,2)"
    assert rs.column_dimensions["A"].width == 31
    assert "D1:E1" in {str(item) for item in rs.merged_cells.ranges}
    assert rs["D1"].value == "Не менять"

    payload["fields"][0]["value"] = 130000
    second = client.post("/api/application-drafts/generate", json=payload)
    assert second.status_code == 201, second.text
    two = second.json()
    assert two["output_artifact_ref"] == one["output_artifact_ref"]
    assert two["output_file_name"] == "working_kp.xlsx"
    assert two["generation_version_no"] == 2
    assert two["output_storage_uri"] != one["output_storage_uri"]
    assert first_path.is_file()
    assert load_workbook(first_path)["КП"]["B2"].value == 125000
    assert (
        load_workbook(tmp_path / two["output_storage_uri"])["КП"]["B2"].value == 130000
    )

    artifact = (
        session.query(DocumentArtifact)
        .filter_by(artifact_ref=one["output_artifact_ref"])
        .one()
    )
    assert artifact.file_name == "working_kp.xlsx"
    assert artifact.current_version == 2

    first_get = client.get(f"/api/application-drafts/{one['generation_id']}")
    assert first_get.status_code == 200
    assert first_get.json()["output_storage_uri"] == one["output_storage_uri"]


def test_arv063_docx_token_fill_preserves_unrelated_structure(
    client, tmp_path: Path, monkeypatch
):
    monkeypatch.setattr(
        "src.modules.application_package_generator.service.load_config",
        lambda: SimpleNamespace(data_dir=str(tmp_path)),
    )
    deal_id = _deal(client, "docx")
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    template = template_dir / "declaration-template.docx"

    document = Document()
    document.add_heading("Декларация участника", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Организация: ")
    token_run = paragraph.add_run("{{COMPANY_NAME}}")
    token_run.italic = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Цена"
    table.cell(0, 1).paragraphs[0].add_run("{{PRICE}}")
    table.cell(1, 0).text = "Не менять"
    table.cell(1, 1).text = "Сохранить"
    document.save(template)

    template_ref = _artifact(
        client, deal_id, tmp_path, "templates/declaration-template.docx"
    )
    response = client.post(
        "/api/application-drafts/generate",
        json={
            "deal_id": deal_id,
            "template_artifact_ref": template_ref,
            "expected_template_version": 1,
            "document_role": "DECLARATION",
            "fields": [
                {
                    "target_locator": "docx:token:{{COMPANY_NAME}}",
                    "value": 'ООО "Арвектум"',
                    "source_type": "DEAL",
                    "source_ref": deal_id,
                },
                {
                    "target_locator": "docx:token:{{PRICE}}",
                    "value": "449 000 ₽",
                    "source_type": "MANUAL_REVIEW",
                    "source_ref": "operator-confirmed-price",
                },
            ],
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["output_file_name"] == "declaration.docx"
    assert payload["output_format"] == "docx"
    assert payload["review_state"] == "DRAFT_REVIEW_ONLY"

    output = Document(tmp_path / payload["output_storage_uri"])
    assert output.paragraphs[0].text == "Декларация участника"
    assert output.paragraphs[1].text == 'Организация: ООО "Арвектум"'
    assert output.paragraphs[1].runs[1].italic is True
    assert output.tables[0].cell(0, 0).text == "Цена"
    assert output.tables[0].cell(0, 1).text == "449 000 ₽"
    assert output.tables[0].cell(1, 0).text == "Не менять"
    assert output.tables[0].cell(1, 1).text == "Сохранить"


def test_arv063_missing_or_unsafe_targets_fail_closed_before_artifact_creation(
    client, session, tmp_path: Path, monkeypatch
):
    monkeypatch.setattr(
        "src.modules.application_package_generator.service.load_config",
        lambda: SimpleNamespace(data_dir=str(tmp_path)),
    )
    deal_id = _deal(client, "fail-closed")
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    template = template_dir / "form2.xlsx"
    workbook = Workbook()
    workbook.active.title = "Форма2"
    workbook.active["A1"] = "unchanged"
    workbook.save(template)
    template_ref = _artifact(client, deal_id, tmp_path, "templates/form2.xlsx")

    before = session.query(DocumentArtifact).count()
    missing_sheet = client.post(
        "/api/application-drafts/generate",
        json={
            "deal_id": deal_id,
            "template_artifact_ref": template_ref,
            "expected_template_version": 1,
            "document_role": "FORM_2",
            "fields": [
                {
                    "target_locator": "xlsx:cell:НетТакогоЛиста!B2",
                    "value": "x",
                    "source_type": "DEAL",
                    "source_ref": deal_id,
                }
            ],
        },
    )
    assert missing_sheet.status_code == 422
    assert session.query(DocumentArtifact).count() == before

    wrong_locator_kind = client.post(
        "/api/application-drafts/generate",
        json={
            "deal_id": deal_id,
            "template_artifact_ref": template_ref,
            "expected_template_version": 1,
            "document_role": "FORM_2",
            "fields": [
                {
                    "target_locator": "docx:token:{{X}}",
                    "value": "x",
                    "source_type": "DEAL",
                    "source_ref": deal_id,
                }
            ],
        },
    )
    assert wrong_locator_kind.status_code == 422
    assert session.query(DocumentArtifact).count() == before


def test_arv063_unsupported_template_and_cross_deal_source_fail_closed(
    client, tmp_path: Path, monkeypatch
):
    monkeypatch.setattr(
        "src.modules.application_package_generator.service.load_config",
        lambda: SimpleNamespace(data_dir=str(tmp_path)),
    )
    deal_a = _deal(client, "A")
    deal_b = _deal(client, "B")
    templates = tmp_path / "templates"
    templates.mkdir()

    pdf = templates / "unsupported.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    pdf_ref = _artifact(client, deal_a, tmp_path, "templates/unsupported.pdf")
    unsupported = client.post(
        "/api/application-drafts/generate",
        json={
            "deal_id": deal_a,
            "template_artifact_ref": pdf_ref,
            "expected_template_version": 1,
            "document_role": "INVENTORY",
            "fields": [
                {
                    "target_locator": "docx:token:{{X}}",
                    "value": "x",
                    "source_type": "DEAL",
                    "source_ref": deal_a,
                }
            ],
        },
    )
    assert unsupported.status_code == 422

    xlsx = templates / "spec.xlsx"
    workbook = Workbook()
    workbook.active.title = "Spec"
    workbook.active["A1"] = "old"
    workbook.save(xlsx)
    template_ref = _artifact(client, deal_a, tmp_path, "templates/spec.xlsx")

    source_doc = templates / "source.docx"
    Document().save(source_doc)
    foreign_artifact = _artifact(client, deal_b, tmp_path, "templates/source.docx")
    cross_deal = client.post(
        "/api/application-drafts/generate",
        json={
            "deal_id": deal_a,
            "template_artifact_ref": template_ref,
            "expected_template_version": 1,
            "document_role": "SPECIFICATION",
            "fields": [
                {
                    "target_locator": "xlsx:cell:Spec!A1",
                    "value": "new",
                    "source_type": "ARTIFACT",
                    "source_ref": foreign_artifact,
                }
            ],
        },
    )
    assert cross_deal.status_code == 422
