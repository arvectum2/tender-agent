"""No cross-tender price inference from authentic but scope-different original sites."""
from __future__ import annotations

import io
import json
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document

from scripts.price01.audit_original_website_cohort import (
    PINS,
    _bid_summary,
    audit_case,
)

REPORT = json.loads((
    Path(__file__).resolve().parents[2]
    / "docs/roadmap/price01_website_cohort_official_2026-10-18.json"
).read_text())


def _protocol(*, bids):
    root = ET.Element("original")
    for id_, admitted, rank, price in bids:
        node = ET.SubElement(root, "applicationInfo")
        for key, value in (
            ("appNumber", id_),
            ("admitted", str(admitted).lower()),
            ("appRating", str(rank) if rank is not None else None),
            ("finalPrice", str(price) if price is not None else None),
        ):
            if value is not None:
                ET.SubElement(node, key).text = value
    return root


def test_pinned_original_registry_scope_categories_and_no_aksay_price_pooling():
    assert REPORT["schema_version"] == "ta-price01-official-website-original-scope-cohort-v1"
    assert REPORT["original_verified_procedures"] == 6
    assert REPORT["original_procedures_with_admitted_final_prices"] == 5
    assert REPORT["original_procedures_without_final_price"] == 1
    assert REPORT["related_but_nonpooled_additional_procedures"] == 2
    assert REPORT["excluded_scope_mismatch_additional_procedures"] == 3
    assert REPORT["new_admitted_same_basket_procurements"] == 0
    assert REPORT["aksay_website_same_basket_count"] == 1
    assert REPORT["profit_or_win_probability_forecast"] is False
    assert REPORT["autonomous_bid_or_payment"] is False
    assert len({x["registry_number"] for x in REPORT["cases"]}) == 6
    assert {x.number for x in PINS} == {x["registry_number"] for x in REPORT["cases"]}
    for row in REPORT["cases"]:
        assert row["accepted_to_aksay_normalized_basket"] is False
        assert row["admitted"] + row["rejected"] == row["total"]
        for key in ("official_zip_sha256", "original_notice_xml_sha256",
                    "original_final_xml_sha256", "original_technical_scope_docx_sha256"):
            assert len(row[key]) == 64
        assert row["source_scope_anchors_verified"]
        assert row["original_notice_subject"]
        assert row["scope_comparison_reason"]
    assert sum(x["total"] for x in REPORT["cases"]) == 140
    assert sum(x["admitted"] for x in REPORT["cases"]) == 124
    assert sum(x["rejected"] for x in REPORT["cases"]) == 16


def test_single_admitted_official_site_not_fake_competitive_distribution():
    dz = next(x for x in REPORT["cases"] if x["registry_number"] == "0132300001726000686")
    assert (dz["total"], dz["admitted"], dz["rejected"]) == (6, 1, 5)
    assert dz["admitted_price_low_rub"] == "249356.00"
    assert dz["admitted_price_high_rub"] == "249356.00"
    assert dz["scope_comparison_state"] == "NOT_COMPARABLE_REGULATED_OFFICIAL_SITE"


def test_extreme_subruble_price_is_not_clamped_or_recommended():
    site = next(x for x in REPORT["cases"] if x["registry_number"] == "0119200000126013647")
    assert site["admitted_price_low_rub"] == "0.01"
    assert site["extreme_sub_ruble_bid_detected"] is True
    assert site["scope_comparison_state"] == "RELATED_SCOPE_REQUIRES_NORMALIZATION"
    assert site["accepted_to_aksay_normalized_basket"] is False


def test_open_competition_final_hides_price_instead_of_using_nmck():
    spb = next(x for x in REPORT["cases"] if x["registry_number"] == "0372200203126000012")
    assert spb["total"] == 10 and spb["admitted"] == 9
    assert spb["source_price_field_state"] == "MISSING_FINAL_PRICE_IN_ORIGINAL_EOK_PROTOCOL"
    assert spb["admitted_price_median_rub"] is None
    assert spb["scope_comparison_state"] == "NOT_COMPARABLE_CUSTOM_CMS"


@pytest.mark.parametrize("rows", [
    [("x", True, 1, "10"), ("x", True, 2, "11")],
    [("x", False, 1, "10"), ("y", True, 2, "11")],
    [("x", True, 2, "10"), ("y", True, 3, "11")],
    [("x", True, 1, "-1")],
    [("x", True, 1, "NaN")],
    [("x", True, None, "10")],
    [("x", True, 1, "10"), ("y", True, 2, None)],
])
def test_untrusted_final_status_price_ranks_or_duplicates_fail_closed(rows):
    with pytest.raises((ValueError, TypeError)):
        _bid_summary(_protocol(bids=rows))


def test_one_competition_has_no_prices_in_final_protocol_is_explicit_state():
    a = _bid_summary(_protocol(bids=[("x", True, 1, None), ("y", True, 2, None)]))
    assert a["admitted"] == 2 and a["admitted_prices_for_individual_original_only"] == 0
    assert a["source_price_field_state"] == "MISSING_FINAL_PRICE_IN_ORIGINAL_EOK_PROTOCOL"
    assert a["admitted_price_low_rub"] is None


def test_synthetic_first_party_zip_and_docx_integrity_pin_without_network(tmp_path):
    number = "0158300034526000388"
    scope = Document()
    scope.add_paragraph("Услуга по разработке веб сайта.")
    scope.add_paragraph("Срок оказания: в течение 14 рабочих дней. Российская CMS.")
    memory = io.BytesIO()
    scope.save(memory)
    scope_bytes = memory.getvalue()
    notice = ET.Element("source")
    ET.SubElement(notice, "purchaseNumber").text = number
    ET.SubElement(notice, "purchaseObjectInfo").text = "Услуга по разработке веб сайта"
    notice_data = ET.tostring(notice, encoding="utf-8")
    protocol = ET.tostring(
        _protocol(bids=[("x", True, 1, "100")]), encoding="utf-8",
    )
    zip_memory = io.BytesIO()
    with ZipFile(zip_memory, "w") as archive:
        archive.writestr(
            f"epNotificationEZK2020_{number}_1_TEST123.xml", notice_data,
        )
        archive.writestr(
            f"epProtocolEZK2020Final_{number}_1_TEST456.xml", protocol,
        )
    zip_bytes = zip_memory.getvalue()
    pin = replace(
        PINS[0],
        zip_sha256=sha256(zip_bytes).hexdigest(),
        final_sha256=sha256(protocol).hexdigest(),
        scope_sha256=sha256(scope_bytes).hexdigest(),
        scope_evidence_phrases=("в течение 14 рабочих дней", "Российская CMS"),
    )
    eis = tmp_path / "eis" / number
    eis.mkdir(parents=True)
    (eis / "documentation-archive.zip").write_bytes(zip_bytes)
    docs = tmp_path / "docs" / number
    docs.mkdir(parents=True)
    original = docs / "technical-scope.docx"
    original.write_bytes(scope_bytes)
    kwargs = {"eis_root": tmp_path / "eis", "technical_scope_root": tmp_path / "docs"}
    row = audit_case(pin, **kwargs)
    assert row["admitted"] == 1 and row["admitted_price_low_rub"] == "100"
    assert row["original_notice_xml_sha256"] == sha256(notice_data).hexdigest()
    assert row["accepted_to_aksay_normalized_basket"] is False

    original.write_bytes(scope_bytes + b"tampering")
    with pytest.raises(ValueError, match="DOCX SHA"):
        audit_case(pin, **kwargs)
    original.write_bytes(scope_bytes)
    (eis / "documentation-archive.zip").write_bytes(zip_bytes + b"tampering")
    with pytest.raises(ValueError, match="ZIP identity"):
        audit_case(pin, **kwargs)
    (eis / "documentation-archive.zip").write_bytes(zip_bytes)
    altered_pin = replace(pin, scope_evidence_phrases=("unknown accepted false requirement",))
    with pytest.raises(ValueError, match="anchor drift"):
        audit_case(altered_pin, **kwargs)


def test_no_local_original_on_ci_is_not_fabricated(tmp_path):
    with pytest.raises(ValueError, match="missing/unsafe"):
        audit_case(PINS[0], eis_root=tmp_path, technical_scope_root=tmp_path)
