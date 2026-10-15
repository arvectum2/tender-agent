"""Source-attested AI development advice; never independent legal gold or SUT.

These read ONLY the historical twenty-case development originals from the
authorised private Mac source root. They never open the four blind locks.
"""
from __future__ import annotations

import hashlib
import os
from pathlib import Path

import pytest
from docx import Document

_DEV_SOURCE_DIR = "ai-expert-20261012/competitive-missing-11-originals-20261014"

# Exact original binary identities attested before AI advice.
_SOURCE_PROOFS = (
    ("0148200000526000028", "attachment-04.docx",
     "b98dd0a2a09ad857f277b9d6da7858d873ee5d6216d66954fb0fb45575416860"),
    ("0148200000526000028", "attachment-02.docx",
     "918ebcea3930147b45d4125497ac8d36eeb20f2109fc5e368e01c79aa5df6d71"),
    ("0149200002326001555", "attachment-04.docx",
     "98e3be9768d9f062f894117fe5e5782ce96a6d314bdb4374f5c00506e44b9ec9"),
    ("0149200002326001555", "attachment-02.docx",
     "810c0693336171845bf7ab7e3685ba098ae48defad756f716161c1e1316c3a05"),
)


def _historical_source_root() -> Path:
    raw = os.getenv("TA_PRIVATE_R4_SOURCE_ROOT")
    if not raw:
        pytest.skip("Real dev originals are private and not available in CI")
    source = Path(raw) / _DEV_SOURCE_DIR
    if not source.is_dir():
        pytest.skip("Historical dev source receipts are not mounted")
    return source


@pytest.fixture(scope="module")
def original_docs():
    root = _historical_source_root()
    result = {}
    for notice, filename, digest in _SOURCE_PROOFS:
        path = root / notice / filename
        if not path.is_file():
            pytest.fail(f"Expected historical development original missing: {notice}/{filename}")
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
        result[notice, filename] = Document(path)
    return result


def _literal(doc, paragraph_no: int, exact: str) -> None:
    assert exact.casefold() in " ".join(doc.paragraphs[paragraph_no - 1].text.split()).casefold()


def test_bidding_rights_and_fstec_application_documents_distinct(original_docs):
    application = original_docs["0148200000526000028", "attachment-04.docx"]
    _literal(application, 35, "SiteSoft.Platform")
    _literal(application, 36, "исключительных прав")
    _literal(application, 37, "лицензионный договор")
    _literal(application, 38, "сублицензионный договор")
    _literal(application, 39, "в составе заявки Исполнитель должен предоставить копию")
    _literal(application, 39, "ФСТЭК России")


def test_customer_side_fstec_audit_cannot_be_conflated_with_bidding_gate(original_docs):
    application = original_docs["0149200002326001555", "attachment-04.docx"]
    contract = original_docs["0149200002326001555", "attachment-02.docx"]
    _literal(application, 23, "пунктом 1 части 1 статьи 31")
    _literal(application, 24, "Наличие у участника неисключительных прав")
    _literal(application, 24, "Фактор")
    _literal(contract, 265, "ООО «Интернет-Фрегат»")
    _literal(contract, 267, "не имеет возможности предоставлять доступ к исходному коду")
    _literal(contract, 706, "Заказчик вправе своими силами или с привлечением третьих лиц")
    _literal(contract, 706, "лицензии ФСТЭК России")
    _literal(contract, 707, "Заказчик вправе своими силами или с привлечением третьих лиц")
    _literal(contract, 709, "Исполнитель обязан")
    # The FSTEC named parties are CUSTOMER / customer-appointed qualified
    # auditors, not a statement that the bidder itself holds such a licence.
    assert "участник" not in contract.paragraphs[705].text.casefold()


def test_cross_document_ip_warranty_and_phased_cashflow_clauses(original_docs):
    contract = original_docs["0149200002326001555", "attachment-02.docx"]
    _literal(contract, 25, "7 рабочих дней")
    _literal(contract, 26, "Авансирование не предусмотрено")
    _literal(contract, 39, "30.04.2026")
    _literal(contract, 40, "10.12.2026")
    _literal(contract, 42, "30.12.2026")
    _literal(contract, 63, "исключительное право")
    _literal(contract, 64, "посредством заключения лицензионного договора")
    _literal(contract, 67, "третьими лицами")
    _literal(contract, 95, "исходных кодов")
    _literal(contract, 96, "свободно распространяемые компиляторы")
    _literal(contract, 112, "12 (двенадцать) месяцев")
    _literal(contract, 118, "не установлено")
    _literal(contract, 119, "не установлено")


def test_previous_development_contract_december_payment_timeline(original_docs):
    contract = original_docs["0148200000526000028", "attachment-02.docx"]
    _literal(contract, 20, "без авансирования")
    _literal(contract, 21, "28 декабря 2027")
    _literal(contract, 39, "10 числа месяца")
    _literal(contract, 47, "20 (Двадцати) рабочих дней")
    assert len(contract.tables) >= 4
    calendar_december = contract.tables[3].rows[12].cells[2].text
    assert "01.12.2027" in calendar_december
    assert "31.12.2027" in calendar_december
